from __future__ import annotations

import io
import unittest
import zipfile
from unittest.mock import patch
from uuid import uuid4

from openpyxl import Workbook

from domains.simple_products.imports.domain import (
    ProductImportTerminalError,
)
from domains.simple_products.imports.infrastructure import parsers
from domains.simple_products.imports.infrastructure.parsers import (
    MAX_IMPORT_ROWS,
    ParsedRow,
    open_source,
)
from domains.simple_products.imports.infrastructure.repository import (
    insert_staged_rows,
)


def _xlsx_payload_with_blank_row() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    sheet.append(
        [
            "Tea",
            "1.000",
        ]
    )
    sheet.append(
        [
            None,
            None,
        ]
    )
    sheet.append(
        [
            "Coffee",
            "2.000",
        ]
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _shared_strings_xlsx_payload() -> bytes:
    parts = {
        "[Content_Types].xml": """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>""",
        "_rels/.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        "xl/workbook.xml": """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
          xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="Sheet1" sheetId="1" r:id="rId1"/>
  </sheets>
</workbook>""",
        "xl/_rels/workbook.xml.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>""",
        "xl/styles.xml": """<?xml version="1.0" encoding="UTF-8"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <numFmts count="0"/>
  <fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>
  <fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>""",
        "xl/sharedStrings.xml": """<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" uniqueCount="3">
  <si><t>Product</t></si>
  <si><t>Unit Price</t></si>
  <si><t>Tea</t></si>
</sst>""",
        "xl/worksheets/sheet1.xml": """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>
    <row r="1">
      <c r="A1" t="s"><v>0</v></c>
      <c r="B1" t="s"><v>1</v></c>
    </row>
    <row r="2">
      <c r="A2" t="s"><v>2</v></c>
      <c r="B2"><v>1.25</v></c>
    </row>
  </sheetData>
</worksheet>""",
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(
        buffer,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        for name, payload in parts.items():
            archive.writestr(
                name,
                payload,
            )
    return buffer.getvalue()


def _csv_payload(
    count: int,
) -> bytes:
    buffer = io.BytesIO()
    buffer.write(
        b"Product,Unit Price\n"
    )
    for index in range(
        count
    ):
        buffer.write(
            (
                f"P{index},1.000\n"
            ).encode("utf-8")
        )
    return buffer.getvalue()


class ParserStreamingTests(
    unittest.TestCase
):
    def test_csv_stream_preserves_original_row_numbers(
        self,
    ) -> None:
        payload = (
            b"Product,Unit Price\n"
            b"Tea,1.000\n"
            b",\n"
            b"Coffee,2.000\n"
        )
        with open_source(
            "products.csv",
            payload,
        ) as source:
            self.assertIs(
                iter(source.rows),
                source.rows,
            )
            rows = list(
                source.rows
            )

        self.assertEqual(
            [
                row.row_number
                for row in rows
            ],
            [2, 4],
        )

    def test_xlsx_stream_preserves_original_row_numbers(
        self,
    ) -> None:
        with open_source(
            "products.xlsx",
            _xlsx_payload_with_blank_row(),
        ) as source:
            rows = list(
                source.rows
            )

        self.assertEqual(
            [
                row.row_number
                for row in rows
            ],
            [2, 4],
        )

    def test_row_limit_is_enforced_incrementally(
        self,
    ) -> None:
        yielded = 0
        with self.assertRaises(
            ProductImportTerminalError
        ):
            with open_source(
                "products.csv",
                _csv_payload(
                    MAX_IMPORT_ROWS
                    + 1
                ),
            ) as source:
                for _row in source.rows:
                    yielded += 1

        self.assertEqual(
            yielded,
            MAX_IMPORT_ROWS,
        )

    def test_xlsx_shared_strings_parse_through_disk_backed_reader(
        self,
    ) -> None:
        with open_source(
            "products.xlsx",
            _shared_strings_xlsx_payload(),
        ) as source:
            rows = list(
                source.rows
            )

        self.assertEqual(
            source.headers,
            [
                "Product",
                "Unit Price",
            ],
        )
        self.assertEqual(
            len(rows),
            1,
        )
        self.assertEqual(
            rows[0].row_number,
            2,
        )
        self.assertEqual(
            rows[0].raw["Product"],
            "Tea",
        )
        self.assertEqual(
            rows[0].raw[
                "Unit Price"
            ],
            "1.25",
        )

    def test_xlsx_security_runs_before_workbook_traversal_and_closes(
        self,
    ) -> None:
        events: list[str] = []

        class FakeCell:
            def __init__(
                self,
                value,
            ) -> None:
                self.value = value
                self.data_type = "s"
                self.number_format = "General"

        class FakeSheet:
            title = "Products"
            sheet_state = "visible"

            def iter_rows(
                self,
                *,
                values_only: bool,
            ):
                self_values = (
                    (
                        FakeCell(
                            "Product"
                        ),
                    ),
                    (
                        FakeCell(
                            "Tea"
                        ),
                    ),
                )
                return iter(
                    self_values
                )

        class FakeWorkbook:
            def __init__(
                self,
            ) -> None:
                self.sheet = FakeSheet()
                self.sheetnames = [
                    "Products"
                ]
                self.worksheets = [
                    self.sheet
                ]
                self.closed = False

            def __getitem__(
                self,
                name: str,
            ):
                if name != "Products":
                    raise KeyError(
                        name
                    )
                return self.sheet

            def close(
                self,
            ) -> None:
                self.closed = True
                events.append(
                    "close"
                )

        class FakeReader:
            def __init__(
                self,
                workbook,
            ) -> None:
                self.wb = workbook

            def close_bounded_resources(
                self,
                *,
                close_archive: bool,
            ) -> None:
                events.append(
                    "close_store"
                )

        workbooks: list[
            FakeWorkbook
        ] = []

        def validate(
            _payload: bytes,
        ) -> None:
            events.append(
                "validate"
            )

        def load(
            *_args,
            **_kwargs,
        ):
            events.append(
                "load"
            )
            workbook = (
                FakeWorkbook()
            )
            workbooks.append(
                workbook
            )
            return FakeReader(
                workbook
            )

        with (
            patch.object(
                parsers,
                "validate_source_content",
                lambda *_args, **_kwargs: None,
            ),
            patch.object(
                parsers,
                "_validate_xlsx_archive",
                validate,
            ),
            patch.object(
                parsers,
                "_load_bounded_workbook",
                load,
            ),
        ):
            with open_source(
                "products.xlsx",
                b"fake",
            ) as source:
                self.assertEqual(
                    list(
                        source.rows
                    )[0].raw[
                        "Product"
                    ],
                    "Tea",
                )

        self.assertEqual(
            events,
            [
                "validate",
                "load",
                "load",
                "close",
                "close_store",
                "close",
                "close_store",
            ],
        )
        self.assertEqual(
            len(
                workbooks
            ),
            2,
        )
        self.assertTrue(
            all(
                workbook.closed
                for workbook
                in workbooks
            )
        )

    def test_xlsx_workbook_closes_when_row_iteration_fails(
        self,
    ) -> None:
        class FakeCell:
            def __init__(
                self,
                value,
            ) -> None:
                self.value = value
                self.data_type = "s"
                self.number_format = "General"

        class BrokenSheet:
            title = "Products"
            sheet_state = "visible"

            def iter_rows(
                self,
                *,
                values_only: bool,
            ):
                yield (
                    FakeCell(
                        "Product"
                    ),
                )
                raise RuntimeError(
                    "read failure"
                )

        class FakeWorkbook:
            def __init__(
                self,
            ) -> None:
                self.sheet = (
                    BrokenSheet()
                )
                self.sheetnames = [
                    "Products"
                ]
                self.worksheets = [
                    self.sheet
                ]
                self.closed = False

            def __getitem__(
                self,
                name: str,
            ):
                if name != "Products":
                    raise KeyError(
                        name
                    )
                return self.sheet

            def close(
                self,
            ) -> None:
                self.closed = True

        class FakeReader:
            def __init__(
                self,
                workbook,
            ) -> None:
                self.wb = workbook

            def close_bounded_resources(
                self,
                *,
                close_archive: bool,
            ) -> None:
                return None

        workbooks: list[
            FakeWorkbook
        ] = []

        def load(
            *_args,
            **_kwargs,
        ):
            workbook = (
                FakeWorkbook()
            )
            workbooks.append(
                workbook
            )
            return FakeReader(
                workbook
            )

        with (
            patch.object(
                parsers,
                "validate_source_content",
                lambda *_args, **_kwargs: None,
            ),
            patch.object(
                parsers,
                "_validate_xlsx_archive",
                lambda _payload: None,
            ),
            patch.object(
                parsers,
                "_load_bounded_workbook",
                load,
            ),
        ):
            with self.assertRaises(
                RuntimeError
            ):
                with open_source(
                    "products.xlsx",
                    b"fake",
                ) as source:
                    list(
                        source.rows
                    )

        self.assertEqual(
            len(
                workbooks
            ),
            2,
        )
        self.assertTrue(
            all(
                workbook.closed
                for workbook
                in workbooks
            )
        )


class _FakeDb:
    def __init__(
        self,
    ) -> None:
        self.batch_sizes: list[
            int
        ] = []

    async def execute(
        self,
        statement,
    ) -> None:
        self.batch_sizes.append(
            len(statement._multi_values[0])
        )


def _parsed_rows(
    count: int,
):
    for index in range(
        count
    ):
        yield ParsedRow(
            row_number=index + 2,
            raw={
                "Product":
                    f"P{index}",
            },
        )


class RepositoryStreamingTests(
    unittest.IsolatedAsyncioTestCase
):
    async def test_staging_batches_are_bounded_at_fifty_thousand_rows(
        self,
    ) -> None:
        db = _FakeDb()
        total = await insert_staged_rows(
            db,
            company_id=1,
            job_id=uuid4(),
            rows=_parsed_rows(
                MAX_IMPORT_ROWS
            ),
            batch_size=1_000,
            max_rows=MAX_IMPORT_ROWS,
        )

        self.assertEqual(
            total,
            MAX_IMPORT_ROWS,
        )
        self.assertEqual(
            len(db.batch_sizes),
            50,
        )
        self.assertEqual(
            max(db.batch_sizes),
            1_000,
        )
        self.assertEqual(
            min(db.batch_sizes),
            1_000,
        )

    async def test_repository_defensively_rejects_row_fifty_thousand_and_one(
        self,
    ) -> None:
        db = _FakeDb()
        with self.assertRaises(
            ProductImportTerminalError
        ):
            await insert_staged_rows(
                db,
                company_id=1,
                job_id=uuid4(),
                rows=_parsed_rows(
                    MAX_IMPORT_ROWS
                    + 1
                ),
                batch_size=1_000,
                max_rows=MAX_IMPORT_ROWS,
            )

        self.assertEqual(
            len(db.batch_sizes),
            50,
        )
        self.assertEqual(
            max(db.batch_sizes),
            1_000,
        )


if __name__ == "__main__":
    unittest.main()
