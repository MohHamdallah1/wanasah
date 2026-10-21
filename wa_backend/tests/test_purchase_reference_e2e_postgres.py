"""Isolated V1 purchase-reference journey on a migrated, disposable PostgreSQL 16 DB.

Each accounting-method scenario uses one company for Product HTTP creation,
reference HTTP edits, real Inbound HTTP commits, concurrency and legacy/import
checks. A second company is a negative RLS fixture only. The sole HTTP auth
override is the authenticated canonical context; real capability/legacy-actor
dependencies, routers, error middleware, domain services and SQL transactions run.
No create_all, SQLite, SAVEPOINT-as-commit harness or mocked financial writer.
"""
from contextlib import asynccontextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
import asyncio
import hashlib
import json
import os
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

from alembic.config import Config as AlembicConfig
from alembic.script import ScriptDirectory
from fastapi import FastAPI, Request
from fastapi.exceptions import RequestValidationError
import httpx
from openpyxl import load_workbook
import pytest
import pytest_asyncio
from sqlalchemy import select, text
from sqlalchemy.engine import make_url
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession, create_async_engine
from sqlalchemy.pool import NullPool
from starlette.exceptions import HTTPException

from api import simple_products
from api.authentication.dependencies import get_current_backoffice_context
from api.warehouse import inbound
from context import tenant_context
from database import get_db
from domains.authorization.capabilities import PERMISSIONS
from domains.inventory_costing.purchase_reference.api import router as reference_router
from domains.inventory_costing.purchase_reference.public import MANAGE, READ
from domains.inventory_supplier_evidence import record_supplier_evidence
from domains.pricing.resolver import resolve_prices_bulk
from domains.simple_products.imports.application.execution_service import build_product_spec
from domains.simple_products.imports.application.purchase_reference_import import record_import_reference_batch
from domains.inventory_costing.purchase_reference.repository import acquire_writer_guard
from domains.simple_products.imports.domain.localization import CANONICAL_IMPORT_FIELDS
from domains.simple_products.imports.domain.mapping import suggest_mapping, validate_mapping
from domains.simple_products.imports.domain.normalization import normalize_raw_row
from domains.simple_products.imports.infrastructure.parsers import open_source
from domains.simple_products.imports.infrastructure.template import build_product_import_template
from domains.simple_products.service import create_products_and_prices
from domains.suppliers.models import Supplier
from domains.suppliers.public import select_active_supplier
from main import (
    WanasahRawASGIMiddleware, custom_http_exception_handler,
    global_exception_handler, validation_exception_handler,
)
from models import (
    Base, Company, Driver, InventoryLocation, ProductBatch, Role,
    UOM, UserLocationAccess, UserRole, role_permissions,
)
from services import (
    apply_inventory_movements_batch, begin_idempotent_operation, complete_idempotent_operation,
)
from domains.inventory_costing.service import build_purchase_cost_input
from tests.canonical_actor_fixtures import seed_company_owner_fixture, seed_ordinary_backoffice_fixture

pytestmark = [
    pytest.mark.skipif(os.getenv("WANASAH_PURCHASE_REFERENCE_E2E") != "1",
                       reason="Explicit disposable GitHub PostgreSQL E2E gate only"),
    pytest.mark.asyncio,
]

EXPECTED_BASE = "ab6307086516483f9615e5b0a7a6c46a6c205110"
CREATE = ("catalog.manage", "catalog.publish", "pricing.manage")
PRICES = ("price_books", "price_book_assignments", "price_publications", "price_book_entries")
ACCOUNTING = ("inventory_movements", "inventory_supplier_evidence", "inventory_balances",
              "inventory_cost_policies", "inventory_cost_events", "inventory_cost_states",
              "inventory_cost_layers", "inventory_cost_allocations", "product_batches")
REFERENCES = ("purchase_cost_references", "purchase_cost_reference_changes", "purchase_cost_reference_commands")
SNAPSHOT_TABLES = (
    "products", "product_variants", "product_uom_conversions", "product_barcodes", "product_locations",
    "operation_idempotency", "domain_audit_events", "transactional_outbox", "system_audit_logs",
    *PRICES, *ACCOUNTING, *REFERENCES,
    *(name for name in Base.metadata.tables if name.startswith("inventory_live_stock_")),
)
TEN_COLUMNS = ("name", "family", "package_uom", "units_per_package", "package_price",
               "unit_price", "unit_barcode", "package_barcode", "lot_control_mode", "expiry_control_mode")


class ForcedRollback(RuntimeError):
    pass


class ObservedSession(AsyncSession):
    """Test-only observation/failure at the actual outer database commit."""
    async def commit(self):
        self.info["commit_calls"] = self.info.get("commit_calls", 0) + 1
        assert self.info["commit_calls"] == 1, "A mutating HTTP command committed more than once"
        callback = self.info.get("before_commit")
        if callback is not None:
            await self.flush()
            await callback(self)
        await super().commit()


@dataclass
class ClientState:
    started: asyncio.Event = field(default_factory=asyncio.Event)
    pids: list[int] = field(default_factory=list)


class Journey:
    def __init__(self, runtime, admin, tenant, foreign, method):
        self.runtime, self.admin = runtime, admin
        self.t, self.foreign, self.method = tenant, foreign, method
        self.stage = "setup"
        self.evidence = []

    def record(self, name, status="PASS", **facts):
        row = dict(stage=name, status=status, method=self.method, database="wanasah_purchase_reference_e2e",
                   company_id=self.t.company_id, **facts)
        self.evidence.append(row)
        print("PURCHASE_REFERENCE_E2E " + json.dumps(row, sort_keys=True), flush=True)
        report = Path(os.environ["WANASAH_E2E_EVIDENCE"])
        report.parent.mkdir(parents=True, exist_ok=True)
        with report.open("a", encoding="utf-8") as output:
            output.write(json.dumps(row, sort_keys=True) + "\n")

    @asynccontextmanager
    async def session(self, *, tenant=None, before_commit=None):
        company = tenant or self.t
        token = tenant_context.set(company.company_id)
        try:
            async with ObservedSession(self.runtime, expire_on_commit=False) as db:
                await db.execute(text("SELECT set_config('app.current_tenant',:c,false)"),
                                 {"c": str(company.company_id)})
                await db.execute(text("SET statement_timeout='15000ms'"))
                await db.execute(text("SET lock_timeout='10000ms'"))
                db.info["before_commit"] = before_commit
                try:
                    yield db
                finally:
                    await db.rollback()
        finally:
            tenant_context.reset(token)

    @asynccontextmanager
    async def client(self, *, actor=None, tenant=None, before_commit=None):
        company = tenant or self.t
        actor = actor or company.owner
        app = FastAPI()
        app.include_router(simple_products.router)
        app.include_router(reference_router)
        app.include_router(inbound.router)
        app.add_exception_handler(HTTPException, custom_http_exception_handler)
        app.add_exception_handler(RequestValidationError, validation_exception_handler)
        app.add_exception_handler(Exception, global_exception_handler)
        app.add_middleware(WanasahRawASGIMiddleware)
        state = ClientState()

        async def database_override(request: Request):
            async with self.session(tenant=company, before_commit=before_commit) as db:
                state.pids.append(await db.scalar(text("SELECT pg_backend_pid()")))
                state.started.set()
                yield db

        async def context_override():
            return actor.context

        app.dependency_overrides[get_db] = database_override
        app.dependency_overrides[get_current_backoffice_context] = context_override
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app, raise_app_exceptions=False),
                                     base_url="http://purchase-reference-e2e", timeout=20) as client:
            yield client, state

    async def send(self, method, path, body=None, **client_options):
        async with self.client(**client_options) as (client, _):
            return await client.request(method, path, json=body)

    async def snapshot(self, tables=SNAPSHOT_TABLES, *, tenant=None):
        async with self.session(tenant=tenant) as db:
            return {table: list((await db.scalars(text(
                f"SELECT to_jsonb(t)::text FROM {table} t WHERE company_id=:c ORDER BY to_jsonb(t)::text"
            ), {"c": (tenant or self.t).company_id})).all()) for table in tables}

    async def grant(self, codes, *, location=None):
        async with AsyncSession(self.admin, expire_on_commit=False) as db:
            await db.execute(text("SELECT set_config('app.current_tenant',:c,true)"), {"c": str(self.t.company_id)})
            role = Role(company_id=self.t.company_id, name="E2E-" + uuid4().hex)
            db.add(role)
            await db.flush()
            rows = list((await db.execute(text(
                "SELECT id,code FROM permissions WHERE code=ANY(:codes)"
            ), {"codes": list(codes)})).all())
            assert {row.code for row in rows} == set(codes)
            await db.execute(role_permissions.insert(), [
                dict(company_id=self.t.company_id, role_id=role.id, permission_id=row.id) for row in rows])
            if location is None:
                db.add(UserRole(company_id=self.t.company_id, role_id=role.id, driver_id=self.t.ordinary.driver_id))
            else:
                db.add(UserLocationAccess(company_id=self.t.company_id, role_id=role.id,
                                         driver_id=self.t.ordinary.driver_id, location_id=location))
            await db.commit()

    async def wait_blocked(self, waiter, blocker):
        # Prove a PostgreSQL lock wait instead of assuming concurrency from a sleep.
        async with self.admin.connect() as db:
            for _ in range(200):
                blockers = await db.scalar(text("SELECT pg_blocking_pids(:pid)"), {"pid": waiter})
                if blocker in blockers:
                    return
                await asyncio.sleep(0.02)
        raise AssertionError(f"Writer {waiter} never waited on transaction {blocker}")


def checked_urls():
    assert os.getenv("GITHUB_ACTIONS") == "true", "This suite may only access its disposable GitHub service"
    assert os.getenv("WANASAH_PURCHASE_REFERENCE_E2E") == "1"
    runtime, admin = [make_url(os.environ[name]) for name in ("DATABASE_URL", "DATABASE_URL_MIGRATION")]
    for url in (runtime, admin):
        assert (url.drivername, url.host, url.port, url.database) == (
            "postgresql+asyncpg", "127.0.0.1", 55473, "wanasah_purchase_reference_e2e")
    assert runtime.username == "purchase_reference_e2e_runtime"
    assert admin.username == "purchase_reference_e2e_admin"
    assert os.environ["WANASAH_E2E_EVIDENCE"]
    return runtime, admin


@pytest_asyncio.fixture
async def journey(request):
    runtime_url, admin_url = checked_urls()
    runtime = create_async_engine(runtime_url, poolclass=NullPool)
    admin = create_async_engine(admin_url, poolclass=NullPool)
    method = request.node.callspec.params["method"]
    try:
        async with runtime.connect() as db:
            role = (await db.execute(text(
                "SELECT current_database(),current_user,rolsuper,rolbypassrls "
                "FROM pg_roles WHERE rolname=current_user"
            ))).one()
            assert tuple(role) == ("wanasah_purchase_reference_e2e", "purchase_reference_e2e_runtime", False, False)
            assert await db.scalar(text("SELECT current_setting('server_version_num')::int")) // 10000 == 16
            for table in REFERENCES:
                flags = (await db.execute(text(
                    "SELECT relrowsecurity,relforcerowsecurity FROM pg_class WHERE oid=to_regclass(:name)"
                ), {"name": table})).one()
                assert tuple(flags) == (True, True)
        async with AsyncSession(admin, expire_on_commit=False) as db:
            cfg = AlembicConfig(str(Path(__file__).resolve().parents[1] / "alembic.ini"))
            assert await db.scalar(text("SELECT version_num FROM alembic_version")) == ScriptDirectory.from_config(cfg).get_current_head()
            for code in PERMISSIONS:
                await db.execute(text("INSERT INTO permissions(code) VALUES(:code) ON CONFLICT(code) DO NOTHING"), {"code": code})
            uoms = dict((await db.execute(select(UOM.code, UOM.id).where(UOM.code.in_(["EACH", "CARTON"])))).all())
            assert set(uoms) == {"EACH", "CARTON"}
            tenants = []
            for index in range(2):
                company = Company(name="Purchase reference E2E " + method, company_code="PCR-E2E-" + uuid4().hex,
                                  currency_code="JOD", timezone="Asia/Amman")
                db.add(company)
                await db.flush()
                await db.execute(text("SELECT set_config('app.current_tenant',:c,true)"), {"c": str(company.id)})
                owner = await seed_company_owner_fixture(db, company.id, password="SyntheticE2EPassword123")
                ordinary = await seed_ordinary_backoffice_fixture(db, company.id)
                wh = InventoryLocation(company_id=company.id, code="E2E-WH", name="Synthetic warehouse",
                                       location_type="WAREHOUSE", is_active=True)
                other = InventoryLocation(company_id=company.id, code="E2E-WH2", name="Other synthetic warehouse",
                                          location_type="WAREHOUSE", is_active=True)
                supplier = Supplier(company_id=company.id, name="Synthetic E2E Supplier",
                                    created_by=owner.driver_id, updated_by=owner.driver_id)
                db.add_all([wh, other, supplier])
                await db.flush()
                tenants.append(SimpleNamespace(company_id=company.id, owner=owner, ordinary=ordinary,
                    warehouse_id=wh.id, other_warehouse_id=other.id, supplier_id=supplier.id,
                    each_id=uoms["EACH"], carton_id=uoms["CARTON"]))
            await db.commit()
        yield Journey(runtime, admin, *tenants, method)
    finally:
        await runtime.dispose()
        await admin.dispose()


def ok(response, status=200):
    assert response.status_code == status, response.text
    UUID(response.headers["x-request-id"])
    return response.json()


def denied(response, status, code=None):
    assert response.status_code == status, response.text
    error = response.json()["error"]
    assert error["request_id"] == response.headers["x-request-id"]
    UUID(error["request_id"])
    assert isinstance(error["context"], dict)
    if code is not None:
        assert error["code"] == code, response.text
    return error


def product(t, *, reference=True, **changes):
    body = dict(request_id=str(uuid4()), name="E2E SKU", family_mode="none", package_uom_code="CARTON",
                units_per_package=12, package_price="60", unit_price=None,
                lot_control_mode="NONE", expiry_control_mode="NONE")
    if reference:
        body["purchase_cost_reference"] = dict(amount="24", purchase_uom_id=t.carton_id, currency_code="JOD")
    body.update(changes)
    return body


def manual(t, revision, amount, *, uom=None, **changes):
    body = dict(request_id=str(uuid4()), expected_revision=revision, amount=amount,
                purchase_uom_id=uom or t.carton_id, currency_code="JOD")
    body.update(changes)
    return body


def receipt(t, variant, cost="48", *, uom=None, quantity="1", **changes):
    rid = str(uuid4())
    body = dict(request_id=rid, supplier_id=t.supplier_id, location_id=t.warehouse_id,
                reference_id="E2E-INVOICE-" + rid,
                items=[dict(product_variant_id=variant, quantity=quantity, uom_id=uom or t.carton_id, unit_cost=cost)])
    body.update(changes)
    return body


async def reference(g, variant):
    return ok(await g.send("GET", f"/purchase-cost-references/{variant}"))["reference"]


async def valuation(g, variant):
    async with g.session() as db:
        return (await db.execute(text(
            "SELECT quantity,inventory_value,average_unit_cost FROM inventory_cost_states "
            "WHERE company_id=:c AND product_variant_id=:v"
        ), {"c": g.t.company_id, "v": variant})).one()


async def price_check(g, variant):
    async with g.session() as db:
        prices = await resolve_prices_bulk(db, company_id=g.t.company_id,
            pairs=[(variant, g.t.each_id), (variant, g.t.carton_id)], customer_id=None, branch_id=None,
            as_of=datetime.now(timezone.utc))
        assert prices[(variant, g.t.each_id)].amount == Decimal("5")
        assert prices[(variant, g.t.carton_id)].amount == Decimal("60")


async def history(g, table, ids):
    async with g.session() as db:
        return dict((await db.execute(text(
            f"SELECT id,to_jsonb(t)::text FROM {table} t WHERE company_id=:c AND id=ANY(:ids) ORDER BY id"
        ), {"c": g.t.company_id, "ids": list(ids)})).all())


async def unchanged_ids(g, table, baseline):
    assert await history(g, table, baseline) == baseline, f"Historical {table} changed"


async def pending_probe(g, before, expected_revision, request_id, *, fail=False):
    called = []

    async def probe(db):
        row = await db.scalar(text(
            "SELECT revision FROM purchase_cost_reference_changes WHERE company_id=:c AND request_id=:rid"
        ), {"c": g.t.company_id, "rid": UUID(request_id)})
        assert row == expected_revision
        completion = (await db.execute(text(
            "SELECT completed_at,response_json FROM operation_idempotency WHERE company_id=:c AND request_id=:rid"
        ), {"c": g.t.company_id, "rid": request_id})).one()
        assert completion.completed_at is not None and completion.response_json is not None
        # Independent connection sees none of the still-uncommitted effects.
        assert await g.snapshot() == before
        called.append(True)
        if fail:
            raise ForcedRollback("E2E forced rollback before the actual outer commit")
    return probe, called


async def external_exit(g, variant, batch_id, quantity):
    async with g.session() as db:
        await apply_inventory_movements_batch(db, company_id=g.t.company_id, performed_by=g.t.owner.driver_id,
            movements=[dict(product_variant_id=variant, batch_id=batch_id, quantity=Decimal(quantity),
                movement_kind="PHYSICAL", reference_type="PURCHASE_REFERENCE_E2E_EXIT",
                reference_id="E2E-COGS-" + uuid4().hex, idempotency_key="E2E-OUT-" + uuid4().hex,
                source_location_id=g.t.warehouse_id, source_stock_status="AVAILABLE",
                destination_location_id=None, destination_stock_status=None)])
        await db.commit()


async def concurrent(g, left_method, left_path, left_body, right_method, right_path, right_body):
    entered, release = asyncio.Event(), asyncio.Event()
    async def hold(db):
        entered.set()
        await asyncio.wait_for(release.wait(), 8)
    async with g.client(before_commit=hold) as (left, ls), g.client() as (right, rs):
        a = asyncio.create_task(left.request(left_method, left_path, json=left_body))
        b = None
        try:
            await asyncio.wait_for(entered.wait(), 8)
            b = asyncio.create_task(right.request(right_method, right_path, json=right_body))
            await asyncio.wait_for(rs.started.wait(), 5)
            await g.wait_blocked(rs.pids[-1], ls.pids[-1])
            assert not b.done()
            release.set()
            return await asyncio.wait_for(a, 12), await asyncio.wait_for(b, 12)
        finally:
            release.set()
            for task in (a, b):
                if task is not None and not task.done():
                    task.cancel()
                    await asyncio.gather(task, return_exceptions=True)


def old_hash(body, *, inbound_items=False):
    # Frozen pre-reference producer: canonical JSON, existing sorted receipt items,
    # excluding only the historical absent Supplier/family-mode fields.
    if inbound_items:
        body = dict(body)
        body["items"] = sorted(body["items"], key=lambda item: json.dumps(
            item, sort_keys=True, separators=(",", ":"), ensure_ascii=False))
    return hashlib.sha256(json.dumps(body, sort_keys=True, separators=(",", ":"), ensure_ascii=False).encode()).hexdigest()


async def legacy_receipt(g, variant, *, supplier_present):
    """Seed genuinely posted pre-feature receipt evidence via the accounting kernel."""
    from schemas import UpgradedInboundRequest
    t = g.t
    payload = receipt(t, variant, supplier_id=t.supplier_id if supplier_present else None, items=[
        dict(product_variant_id=variant, quantity="1", uom_id=t.each_id, unit_cost="2"),
        dict(product_variant_id=variant, quantity="1", uom_id=t.carton_id, unit_cost="24"),
    ])
    model = UpgradedInboundRequest(**payload)
    # Spell out the historical shape; never call the new receipt hash helper.
    dumped = model.model_dump(mode="json")
    body = {key: dumped[key] for key in ("supplier_id", "location_id", "reference_id", "notes", "items")}
    if not supplier_present:
        body.pop("supplier_id")
    digest = old_hash(body, inbound_items=True)
    async with g.session() as db:
        op, replay = await begin_idempotent_operation(db, company_id=t.company_id,
            actor_id=t.owner.driver_id, operation="WAREHOUSE_INBOUND",
            request_id=payload["request_id"], request_hash=digest)
        assert replay is None
        movements = []
        for item in model.items:
            batch = ProductBatch(company_id=t.company_id, product_variant_id=variant,
                                 batch_number="E2E-LEGACY-" + uuid4().hex)
            db.add(batch)
            await db.flush()
            factor = Decimal(12 if item.uom_id == t.carton_id else 1)
            qty = item.quantity * factor
            movements.extend(await apply_inventory_movements_batch(db, company_id=t.company_id,
                performed_by=t.owner.driver_id, movements=[dict(
                    product_variant_id=variant, batch_id=batch.id, quantity=qty,
                    movement_kind="PHYSICAL", reference_type="INBOUND_SUPPLIER",
                    reference_id=payload["reference_id"], idempotency_key="LEGACY-" + uuid4().hex,
                    source_location_id=None, source_stock_status=None,
                    destination_location_id=t.warehouse_id, destination_stock_status="AVAILABLE",
                    inventory_cost_input=build_purchase_cost_input(input_uom_id=item.uom_id,
                        input_quantity=item.quantity, input_unit_cost=item.unit_cost, base_quantity=qty))]))
        if supplier_present:
            supplier = await select_active_supplier(db, access=t.owner.access(db), supplier_id=t.supplier_id)
            record_supplier_evidence(db, company_id=t.company_id, movement_ids=[row.id for row in movements],
                                     supplier=supplier, request_id=UUID(payload["request_id"]))
        complete_idempotent_operation(op, {"message": "INBOUND_POSTED"})
        await db.commit()
    if not supplier_present:
        payload.pop("supplier_id")
    before = await g.snapshot()
    ok(await g.send("POST", "/warehouse/inbound", payload), 201)
    assert await g.snapshot() == before


async def import_contract(g, original_variant):
    assert tuple(CANONICAL_IMPORT_FIELDS[:10]) == TEN_COLUMNS
    assert CANONICAL_IMPORT_FIELDS[10] == "purchase_unit_cost"
    for locale in ("ar", "en"):
        blob = build_product_import_template(locale=locale)
        wb = load_workbook(BytesIO(blob))
        sheet = wb.active
        assert sheet.max_column == 11 and sheet.auto_filter.ref.startswith("A1:K")
        sheet.append(["E2E imported " + locale, None, "CARTON", "12", "60", None,
                      None, None, "NONE", "NONE", "1.250000"])
        output = BytesIO()
        wb.save(output)
        wb.close()
        with open_source("e2e-" + locale + ".xlsx", output.getvalue()) as source:
            assert len(source.headers) == 11
            mapping = validate_mapping(source.headers, suggest_mapping(source.headers))
            assert set(mapping) == set(TEN_COLUMNS) | {"purchase_unit_cost"}
            rows = list(source.rows)
            assert len(rows) == 1
            data = normalize_raw_row(rows[0].raw, mapping,
                default_lot_control_mode="NONE", default_expiry_control_mode="NONE")
            assert data["purchase_unit_cost"] == "1.250000"
            with pytest.raises(ValueError, match="Unknown mapping field"):
                validate_mapping(source.headers, {**mapping, "purchase_cost_reference": source.headers[0]})
        before = await g.snapshot(ACCOUNTING)
        # True Product + reference + audit share one tenant-scoped database commit.
        # Worker queue orchestration is tested separately from this full DB boundary.
        async with g.session() as db:
            actor = await db.get(Driver, g.t.owner.driver_id)
            request_id = uuid4()
            await acquire_writer_guard(db, company_id=g.t.company_id)
            created = await create_products_and_prices(db, actor=actor, company_id=g.t.company_id,
                request_id=request_id, specs=[build_product_spec(data)])
            imported_id = created[0][0].id
            await record_import_reference_batch(
                db, access=g.t.owner.access(db), company_id=g.t.company_id,
                job_id=uuid4(), request_id=request_id, request_hash=hashlib.sha256(
                    str(request_id).encode()).hexdigest(),
                rows=[SimpleNamespace(id=1, normalized_data=data)], created=created,
            )
            await db.commit()
        assert await g.snapshot(ACCOUNTING) == before
        assert (await reference(g, imported_id))["amount"] == "1.250000"
        async with g.session() as db:
            row = (await db.execute(text(
                "SELECT source,purchase_uom_id,revision FROM purchase_cost_references "
                "WHERE company_id=:c AND product_variant_id=:v"
            ), {"c": g.t.company_id, "v": imported_id})).one()
            assert row.source == "PRODUCT_IMPORT" and row.purchase_uom_id == g.t.each_id and row.revision == 1
            assert await db.scalar(text(
                "SELECT count(*) FROM purchase_cost_reference_changes "
                "WHERE company_id=:c AND product_variant_id=:v"
            ), {"c": g.t.company_id, "v": imported_id}) == 1
        await price_check(g, imported_id)
        await price_check(g, original_variant)


async def complete_journey(g):
    t = g.t
    g.stage = "01_product_create_atomic"
    ok(await g.send("PUT", "/warehouse/costing-policy",
                   dict(request_id=str(uuid4()), method=g.method, expected_version=0)))
    before = await g.snapshot()
    create = product(t)
    probe, calls = await pending_probe(g, before, 1, create["request_id"])
    created = ok(await g.send("POST", "/simple-products", create, before_commit=probe), 201)
    assert calls == [True]
    v = created["product_variant_id"]
    first = await reference(g, v)
    assert (first["amount"], first["currency_code"], first["purchase_uom_id"], first["revision"]) == (
        "24.000000", "JOD", t.carton_id, 1)
    assert first["provenance"]["reason"] == "PRODUCT_CREATE"
    g.record(g.stage, variant_id=v, reference_revision=1, independent_commit_visibility=True)

    g.stage = "02_read_manual_revision"
    before = await g.snapshot(ACCOUNTING + PRICES)
    quote = ok(await g.send("GET", f"/purchase-cost-references/{v}?quote_uom_id={t.each_id}"))["reference"]
    assert quote["quote"]["amount"] == "2.000000"
    changed = ok(await g.send("PUT", f"/purchase-cost-references/{v}", manual(t, 1, "36")))["reference"]
    assert changed["revision"] == 2
    async with g.session() as db:
        assert await db.scalar(text("SELECT version FROM product_variants WHERE company_id=:c AND id=:v"),
                               {"c": t.company_id, "v": v}) == 2
    assert await g.snapshot(ACCOUNTING + PRICES) == before
    g.record(g.stage, variant_id=v, catalog_version=2, reference_revision=2, exact_each_quote="2.000000")

    g.stage = "03_actual_receipt_atomic_provenance"
    options = ok(await g.send("POST", "/warehouse/inbound/options", {"ids": [v]}))
    assert options["items"][0]["purchase_cost_reference"]["amount"] == "36.000000"
    incoming = receipt(t, v, "48")
    before = await g.snapshot()
    probe, calls = await pending_probe(g, before, 3, incoming["request_id"])
    received = ok(await g.send("POST", "/warehouse/inbound", incoming, before_commit=probe), 201)
    assert calls == [True]
    ref = await reference(g, v)
    assert (ref["amount"], ref["revision"], ref["source"]) == ("48.000000", 3, "SUPPLIER_INBOUND")
    async with g.session() as db:
        source = (await db.execute(text(
            "SELECT e.company_id,e.product_variant_id,e.input_uom_id,e.input_unit_cost,"
            "m.reference_id,m.batch_id,s.request_id FROM inventory_cost_events e "
            "JOIN inventory_movements m ON m.company_id=e.company_id AND m.id=e.inventory_movement_id "
            "JOIN inventory_supplier_evidence s ON s.company_id=m.company_id AND s.movement_id=m.id "
            "WHERE e.company_id=:c AND e.id=:id"
        ), {"c": t.company_id, "id": ref["source_cost_event_id"]})).one()
        assert (source.company_id, source.product_variant_id, source.input_uom_id, source.input_unit_cost,
                source.reference_id, str(source.request_id)) == (
            t.company_id, v, t.carton_id, Decimal("48"), incoming["reference_id"], incoming["request_id"])
        event_ids = list((await db.scalars(text("SELECT id FROM inventory_cost_events WHERE company_id=:c"),
                                          {"c": t.company_id})).all())
        movement_ids = list((await db.scalars(text("SELECT id FROM inventory_movements WHERE company_id=:c"),
                                             {"c": t.company_id})).all())
        price_ids = list((await db.scalars(text("SELECT id FROM price_book_entries WHERE company_id=:c"),
                                          {"c": t.company_id})).all())
        actor = (await db.execute(text(
            "SELECT actor_principal_id,actor_backoffice_user_id FROM purchase_cost_reference_commands "
            "WHERE company_id=:c AND request_id=:rid"
        ), {"c": t.company_id, "rid": UUID(incoming["request_id"])})).one()
        assert tuple(actor) == (t.owner.context.principal_id, t.owner.context.backoffice_user_id)
        assert actor.actor_principal_id != actor.actor_backoffice_user_id != t.owner.driver_id
    assert tuple(await valuation(g, v)) == (Decimal("12"), Decimal("48"), Decimal("4"))
    for table in PRICES:
        assert (await g.snapshot(PRICES))[table] == before[table]
    invoice = await history(g, "inventory_movements", movement_ids)
    frozen_prices = await history(g, "price_book_entries", price_ids)
    g.record(g.stage, variant_id=v, actual_cost="48.000000", previous_reference="36.000000",
             reference_revision=3, cost_event_id=ref["source_cost_event_id"], invoice=incoming["reference_id"])

    g.stage = "04_historical_cogs"
    await external_exit(g, v, source.batch_id, "3")
    assert tuple(await valuation(g, v)) == (Decimal("9"), Decimal("36"), Decimal("4"))
    async with g.session() as db:
        event_ids = list((await db.scalars(text("SELECT id FROM inventory_cost_events WHERE company_id=:c"),
                                          {"c": t.company_id})).all())
        allocations = (await db.execute(text(
            "SELECT a.quantity,l.unit_cost,a.amount FROM inventory_cost_allocations a "
            "JOIN inventory_cost_layers l ON l.company_id=a.company_id AND l.id=a.cost_layer_id WHERE a.company_id=:c"
        ), {"c": t.company_id})).all()
        outbound = (await db.execute(text(
            "SELECT unit_cost,total_cost FROM inventory_cost_events WHERE company_id=:c AND event_type='OUTBOUND'"
        ), {"c": t.company_id})).one()
        assert tuple(outbound) == (Decimal("4"), Decimal("12"))
        assert len(allocations) == (1 if g.method == "FIFO" else 0)
    immutable_events = await history(g, "inventory_cost_events", event_ids)
    cogs_allocations = (await g.snapshot(("inventory_cost_allocations",)))["inventory_cost_allocations"]
    g.record(g.stage, cogs="12.000000", quantity="9.000000", valuation="36.000000")

    g.stage = "05_old_receipt_replay_after_new_manual"
    before = await g.snapshot(ACCOUNTING + PRICES)
    assert ok(await g.send("PUT", f"/purchase-cost-references/{v}", manual(t, 3, "60")))["reference"]["revision"] == 4
    assert await g.snapshot(ACCOUNTING + PRICES) == before
    committed = await g.snapshot()
    assert ok(await g.send("POST", "/warehouse/inbound", incoming), 201) == received
    assert ok(await g.send("POST", "/simple-products", create), 201) == created
    assert await g.snapshot() == committed
    tampered = json.loads(json.dumps(incoming))
    tampered["items"][0]["unit_cost"] = "49"
    denied(await g.send("POST", "/warehouse/inbound", tampered), 409)
    assert await g.snapshot() == committed
    g.record(g.stage, newer_reference="60.000000", reference_revision=4, whole_company_snapshot_unchanged=True)

    g.stage = "06_mixed_native_cost_uom_selection"
    mixed = receipt(t, v, items=[
        dict(product_variant_id=v, quantity="1", uom_id=t.carton_id, unit_cost="24"),
        dict(product_variant_id=v, quantity="2", uom_id=t.each_id, unit_cost="3"),
    ])
    before = await g.snapshot()
    denied(await g.send("POST", "/warehouse/inbound", mixed), 422, "INBOUND_REFERENCE_SELECTION_REQUIRED")
    assert await g.snapshot() == before
    mixed["purchase_reference_selections"] = [dict(product_variant_id=v, uom_id=t.each_id, unit_cost="3")]
    mixed["items"].reverse()
    ok(await g.send("POST", "/warehouse/inbound", mixed), 201)
    ref = await reference(g, v)
    assert (ref["amount"], ref["purchase_uom_id"], ref["revision"]) == ("3.000000", t.each_id, 5)
    assert ref["provenance"]["selection"] == "EXPLICIT"
    assert str(ref["provenance"]["request_id"]) == mixed["request_id"]
    assert tuple(await valuation(g, v)) == (Decimal("23"), Decimal("66"), Decimal("2.869565"))
    after = await g.snapshot()
    mixed["items"].reverse()
    ok(await g.send("POST", "/warehouse/inbound", mixed), 201)
    assert await g.snapshot() == after
    changed_selection = json.loads(json.dumps(mixed))
    changed_selection["purchase_reference_selections"][0].update(uom_id=t.carton_id, unit_cost="24")
    denied(await g.send("POST", "/warehouse/inbound", changed_selection), 409)
    assert await g.snapshot() == after
    await unchanged_ids(g, "inventory_cost_events", immutable_events)
    await unchanged_ids(g, "inventory_movements", invoice)
    assert (await g.snapshot(("inventory_cost_allocations",)))["inventory_cost_allocations"] == cogs_allocations
    g.record(g.stage, selected_amount="3.000000", selected_uom=t.each_id, reference_revision=5,
             valuation="66.000000", quantity="23.000000", acquisition_costs_preserved=True)

    g.stage = "07_concurrent_manual_receipt"
    left, right = await concurrent(g, "PUT", f"/purchase-cost-references/{v}", manual(t, 5, "84"),
                                  "POST", "/warehouse/inbound", receipt(t, v, "96"))
    assert ok(left)["reference"]["revision"] == 6
    ok(right, 201)
    assert (await reference(g, v))["revision"] == 7
    assert (await reference(g, v))["amount"] == "96.000000"
    assert tuple(await valuation(g, v)) == (Decimal("35"), Decimal("162"), Decimal("4.628571"))
    await unchanged_ids(g, "inventory_cost_events", immutable_events)
    g.record(g.stage, postgres_lock_wait_proven=True, serial_revisions=[6, 7], last_committed_reference="96.000000")

    g.stage = "08_first_insert_race"
    empty = ok(await g.send("POST", "/simple-products", product(t, reference=False, name="E2E first insertion")), 201)
    v2 = empty["product_variant_id"]
    assert await reference(g, v2) is None
    before = await g.snapshot(ACCOUNTING + PRICES)
    left, right = await concurrent(g, "PUT", f"/purchase-cost-references/{v2}", manual(t, 0, "10", uom=t.each_id),
                                  "PUT", f"/purchase-cost-references/{v2}", manual(t, 0, "11", uom=t.each_id))
    assert ok(left)["reference"]["revision"] == 1
    denied(right, 409, "PURCHASE_REFERENCE_REVISION_CONFLICT")
    assert (await reference(g, v2))["amount"] == "10.000000"
    assert await g.snapshot(ACCOUNTING + PRICES) == before
    async with g.session() as db:
        assert await db.scalar(text(
            "SELECT count(*) FROM purchase_cost_reference_changes WHERE company_id=:c AND product_variant_id=:v"
        ), {"c": t.company_id, "v": v2}) == 1
    g.record(g.stage, variant_id=v2, winner_revision=1, loser_status=409, postgres_lock_wait_proven=True)

    g.stage = "09_forced_product_receipt_rollbacks"
    before = await g.snapshot()
    failed_product = product(t, name="E2E rolled back Product")
    probe, calls = await pending_probe(g, before, 1, failed_product["request_id"], fail=True)
    denied(await g.send("POST", "/simple-products", failed_product, before_commit=probe), 500)
    assert calls == [True] and await g.snapshot() == before
    good = ok(await g.send("POST", "/simple-products", failed_product), 201)
    assert (await reference(g, good["product_variant_id"]))["revision"] == 1
    failed_receipt = receipt(t, v2, "12", uom=t.each_id)
    before = await g.snapshot()
    probe, calls = await pending_probe(g, before, 2, failed_receipt["request_id"], fail=True)
    denied(await g.send("POST", "/warehouse/inbound", failed_receipt, before_commit=probe), 500, "INBOUND_INTERNAL_ERROR")
    assert calls == [True] and await g.snapshot() == before
    ok(await g.send("POST", "/warehouse/inbound", failed_receipt), 201)
    assert (await reference(g, v2))["revision"] == 2
    g.record(g.stage, product_and_receipt_rolled_back=True, same_request_retry_succeeded=True)

    g.stage = "10_permissions_tenant_isolation"
    await g.grant(CREATE)
    before = await g.snapshot()
    denied(await g.send("POST", "/simple-products", product(t), actor=t.ordinary), 403, "PURCHASE_REFERENCE_ACCESS_DENIED")
    denied(await g.send("GET", f"/purchase-cost-references/{v}", actor=t.ordinary), 403, "PURCHASE_REFERENCE_ACCESS_DENIED")
    denied(await g.send("PUT", f"/purchase-cost-references/{v}", manual(t, 7, "1"), actor=t.ordinary),
           403, "PURCHASE_REFERENCE_ACCESS_DENIED")
    assert await g.snapshot() == before
    await g.grant((READ, "supplier.read"))
    ok(await g.send("GET", f"/purchase-cost-references/{v}", actor=t.ordinary))
    await g.grant(("inbound.create", MANAGE), location=t.warehouse_id)
    before = await g.snapshot()
    denied(await g.send("PUT", f"/purchase-cost-references/{v}", manual(t, 7, "1"), actor=t.ordinary),
           403, "PURCHASE_REFERENCE_ACCESS_DENIED")
    wrong_location = receipt(t, v, location_id=t.other_warehouse_id)
    denied(await g.send("POST", "/warehouse/inbound", wrong_location, actor=t.ordinary), 403)
    assert await g.snapshot() == before
    ordinary_receipt = receipt(t, v, "6", uom=t.each_id)
    ok(await g.send("POST", "/warehouse/inbound", ordinary_receipt, actor=t.ordinary), 201)
    assert (await reference(g, v))["amount"] == "6.000000"
    async with g.session(tenant=g.foreign) as db:
        for table in REFERENCES:
            assert await db.scalar(text(f"SELECT count(*) FROM {table}")) == 0
    denied(await g.send("GET", f"/purchase-cost-references/{v}", tenant=g.foreign), 404)
    foreign_receipt = receipt(g.foreign, v)
    denied(await g.send("POST", "/warehouse/inbound", foreign_receipt, tenant=g.foreign), 409)
    g.record(g.stage, catalog_pricing_do_not_grant_cost=True, location_manage_denied=True,
             inbound_side_effect_only=True, tenant_isolation=True)

    g.stage = "11_immutable_reference_audit"
    before = await g.snapshot()
    async with g.session() as db:
        for table in ("purchase_cost_reference_changes", "purchase_cost_reference_commands"):
            for action in ("UPDATE " + table + " SET company_id=company_id", "DELETE FROM " + table):
                with pytest.raises(DBAPIError):
                    async with db.begin_nested():
                        await db.execute(text(action))
    assert await g.snapshot() == before
    async with AsyncSession(g.admin) as db:
        with pytest.raises(DBAPIError):
            await db.execute(text("UPDATE purchase_cost_reference_changes SET amount=amount WHERE company_id=:c"),
                             {"c": t.company_id})
    g.record(g.stage, runtime_update_delete_rejected=True, privileged_history_rewrite_rejected=True)

    g.stage = "12_legacy_product_inbound_replay"
    for mode in (None, "none"):
        old_product = product(t, reference=False, family_mode=mode, name="E2E legacy Product")
        model = simple_products.SimpleProductCreate(**old_product)
        dumped = model.model_dump(mode="json")
        keys = ("name", "family_id", "family_name", "family_mode", "package_uom_code", "units_per_package",
                "package_price", "unit_price", "unit_barcode", "package_barcode", "lot_control_mode", "expiry_control_mode")
        old_body = {key: dumped[key] for key in keys}
        if mode is None:
            old_body.pop("family_mode")
        digest = old_hash(old_body)
        result = ok(await g.send("POST", "/simple-products", old_product), 201)
        async with g.session() as db:
            assert await db.scalar(text(
                "SELECT request_hash FROM operation_idempotency WHERE company_id=:c AND operation='SIMPLE_PRODUCT_CREATE_V3' AND request_id=:rid"
            ), {"c": t.company_id, "rid": old_product["request_id"]}) == digest
        lv = result["product_variant_id"]
        assert await reference(g, lv) is None
        ok(await g.send("PUT", f"/purchase-cost-references/{lv}", manual(t, 0, "18", uom=t.each_id)))
        before = await g.snapshot()
        assert ok(await g.send("POST", "/simple-products", old_product), 201) == result
        assert await g.snapshot() == before
    for supplied in (False, True):
        await legacy_receipt(g, v, supplier_present=supplied)
    g.record(g.stage, legacy_product_modes=[None, "none"], legacy_supplier_presence=[False, True],
             ambiguous_legacy_lines_replayed_without_reference_write=True)

    g.stage = "13_ten_column_excel_contract"
    await import_contract(g, v)
    await unchanged_ids(g, "inventory_cost_events", immutable_events)
    await unchanged_ids(g, "inventory_movements", invoice)
    await unchanged_ids(g, "price_book_entries", frozen_prices)
    await price_check(g, v)
    # Reconcile authoritative stock, receipt-cost ledger, COGS and FIFO layers
    # across the whole company after all scenarios, not only empty-table snapshots.
    async with g.session() as db:
        states = (await db.execute(text(
            "SELECT product_variant_id,quantity,inventory_value,average_unit_cost "
            "FROM inventory_cost_states WHERE company_id=:c"
        ), {"c": t.company_id})).all()
        assert len(states) == 2
        for state in states:
            balances = await db.scalar(text(
                "SELECT coalesce(sum(on_hand_quantity),0) FROM inventory_balances WHERE company_id=:c AND product_variant_id=:v"
            ), {"c": t.company_id, "v": state.product_variant_id})
            ledger = (await db.execute(text(
                "SELECT sum(CASE WHEN event_type='OUTBOUND' THEN -quantity ELSE quantity END) AS quantity,"
                "sum(CASE WHEN event_type='OUTBOUND' THEN -total_cost ELSE total_cost END) AS value "
                "FROM inventory_cost_events WHERE company_id=:c AND product_variant_id=:v"
            ), {"c": t.company_id, "v": state.product_variant_id})).one()
            assert state.quantity == balances == ledger.quantity
            assert state.inventory_value == ledger.value
            assert state.average_unit_cost == (state.inventory_value / state.quantity).quantize(
                Decimal("0.000001"), rounding=ROUND_HALF_UP)
            if g.method == "FIFO":
                layers = (await db.execute(text(
                    "SELECT sum(remaining_quantity) AS quantity,sum(remaining_value) AS value "
                    "FROM inventory_cost_layers WHERE company_id=:c AND product_variant_id=:v"
                ), {"c": t.company_id, "v": state.product_variant_id})).one()
                assert (layers.quantity, layers.value) == (state.quantity, state.inventory_value)
        if g.method == "MOVING_AVERAGE":
            assert await db.scalar(text("SELECT count(*) FROM inventory_cost_layers WHERE company_id=:c"),
                                   {"c": t.company_id}) == 0
        assert sum(row.inventory_value for row in states) == Decimal("232")
        assert await db.scalar(text(
            "SELECT sum(total_cost) FROM inventory_cost_events WHERE company_id=:c AND event_type='OUTBOUND'"
        ), {"c": t.company_id}) == Decimal("12")
        policy = (await db.execute(text(
            "SELECT method,is_active,locked_at FROM inventory_cost_policies WHERE company_id=:c"
        ), {"c": t.company_id})).one()
        assert policy.method == g.method and policy.is_active and policy.locked_at is not None
    g.record(g.stage, locales=["ar", "en"], old_columns_preserved=10, optional_column_count=11,
              import_reference_backfill=False,
             company_inventory_value="232.000000", historical_cogs="12.000000",
             old_invoices_cogs_prices_unchanged=True)
    g.record("SCENARIO_COMPLETE", original_variant_id=v, phases_passed=13)


@pytest.mark.parametrize("method", ["FIFO", "MOVING_AVERAGE"])
async def test_v1_purchase_reference_end_to_end(journey, method):
    try:
        await complete_journey(journey)
    except BaseException as exc:
        journey.record(journey.stage, status="FAIL", failure_type=type(exc).__name__, reason=str(exc)[:1000])
        raise
