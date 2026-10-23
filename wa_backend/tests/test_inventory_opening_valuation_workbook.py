import unittest
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from domains.inventory_costing.opening_valuation import (
    OpeningBatchSnapshot,
    OpeningSnapshot,
    OpeningVariantSnapshot,
    build_opening_valuation_template,
    parse_opening_valuation_template,
)
from domains.inventory_costing.service import CostingError


class OpeningValuationWorkbookContractTests(unittest.TestCase):
    def snapshot(self):
        return OpeningSnapshot(
            hash="a" * 64,
            total_quantity=Decimal("6.5"),
            variants=(
                OpeningVariantSnapshot(
                    product_variant_id=1,
                    sku="SKU-1",
                    name="منتج تجريبي",
                    base_uom_id=1,
                    base_uom_code="EACH",
                    base_uom_name="وحدة",
                    quantity=Decimal("6.5"),
                    batches=(
                        OpeningBatchSnapshot(1, 10, Decimal("3.5")),
                        OpeningBatchSnapshot(1, 11, Decimal("3")),
                    ),
                ),
            ),
        )

    def test_template_is_read_only_except_opening_cost_and_round_trips(self):
        payload = build_opening_valuation_template(
            snapshot=self.snapshot(),
            currency_code="JOD",
            locale="ar",
        )
        wb = load_workbook(BytesIO(payload))
        ws = wb["Opening Valuation"]
        self.assertTrue(ws["A2"].protection.locked)
        self.assertTrue(ws["D2"].protection.locked)
        self.assertFalse(ws["E2"].protection.locked)
        self.assertEqual(wb["_wanasah_meta"]["B2"].value, "a" * 64)
        ws["E2"] = Decimal("2.500000")
        output = BytesIO()
        wb.save(output)
        wb.close()

        snapshot_hash, lines = parse_opening_valuation_template(output.getvalue())
        self.assertEqual(snapshot_hash, "a" * 64)
        self.assertEqual(lines, [{"sku": "SKU-1", "unit_cost": 2.5}])

    def test_formula_cost_is_rejected(self):
        payload = build_opening_valuation_template(
            snapshot=self.snapshot(),
            currency_code="JOD",
        )
        wb = load_workbook(BytesIO(payload))
        wb["Opening Valuation"]["E2"] = "=1+1"
        output = BytesIO()
        wb.save(output)
        wb.close()
        with self.assertRaises(CostingError) as caught:
            parse_opening_valuation_template(output.getvalue())
        self.assertEqual(caught.exception.code, "OPENING_VALUATION_FORMULA_NOT_ALLOWED")


if __name__ == "__main__":
    unittest.main()
