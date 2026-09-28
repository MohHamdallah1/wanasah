from __future__ import annotations

import io
import unittest
import zipfile
from datetime import date, datetime, timezone
from xml.etree import ElementTree
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from domains.simple_products.imports.domain import (
    CANONICAL_IMPORT_FIELDS,
    EN_IMPORT_LOCALE,
    ProductImportTerminalError,
    SOURCE_CELL_META_KEY,
    WANASAH_TEMPLATE_MARKER,
    WANASAH_TEMPLATE_META_SHEET,
    normalize_spreadsheet_date,
)
from domains.simple_products.imports.domain.normalization import (
    normalize_raw_row,
)
from domains.simple_products.imports.infrastructure.parsers import (
    open_source,
)
from domains.simple_products.imports.infrastructure.repository import (
    insert_staged_rows,
)
from domains.simple_products.imports.infrastructure.template import (
    MAX_TEMPLATE_ROWS,
    build_product_import_template,
)
from domains.simple_products.service import (
    SimpleProductError,
)


_MAIN_NS = (
    "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
)


def _save_workbook(
    workbook: Workbook,
) -> bytes:
    output = io.BytesIO()
    workbook.save(
        output
    )
    workbook.close()
    return output.getvalue()


def _xlsx_with_blank_row() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    sheet.append(
        [
            "Tea",
            "1.000",
        ]
    )
    sheet.append(
        [
            None,
            None,
        ]
    )
    sheet.append(
        [
            "Coffee",
            "2.000",
        ]
    )
    return _save_workbook(
        workbook
    )


def _external_multisheet_xlsx() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Sheet A"
    first.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    first.append(
        [
            "Tea",
            "1.000",
        ]
    )
    second = workbook.create_sheet(
        "Sheet B"
    )
    second.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    second.append(
        [
            "Coffee",
            "2.000",
        ]
    )
    return _save_workbook(
        workbook
    )


def _xlsx_formula_payload(
    *,
    cached_value: str | None,
    formula: str = "1+1",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    sheet.append(
        [
            "Tea",
            f"={formula}",
        ]
    )
    payload = _save_workbook(
        workbook
    )
    if cached_value is None:
        return payload

    source = io.BytesIO(
        payload
    )
    destination = io.BytesIO()
    with (
        zipfile.ZipFile(
            source,
            "r",
        ) as incoming,
        zipfile.ZipFile(
            destination,
            "w",
            compression=
                zipfile.ZIP_DEFLATED,
        ) as outgoing,
    ):
        for info in incoming.infolist():
            data = incoming.read(
                info.filename
            )
            if (
                info.filename
                == "xl/worksheets/sheet1.xml"
            ):
                root = ElementTree.fromstring(
                    data
                )
                cell = root.find(
                    (
                        ".//"
                        f"{{{_MAIN_NS}}}c"
                        "[@r='B2']"
                    )
                )
                if cell is None:
                    raise AssertionError(
                        "Formula fixture cell B2 was not written."
                    )
                value_node = cell.find(
                    f"{{{_MAIN_NS}}}v"
                )
                if value_node is None:
                    value_node = (
                        ElementTree.SubElement(
                            cell,
                            f"{{{_MAIN_NS}}}v",
                        )
                    )
                value_node.text = str(
                    cached_value
                )
                data = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            outgoing.writestr(
                info,
                data,
            )
    return destination.getvalue()


def _barcode_xlsx(
    value,
    *,
    as_text: bool,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(
        [
            "Product",
            "Unit Price",
            "Unit Barcode",
        ]
    )
    sheet.append(
        [
            "Tea",
            "1.000",
            value,
        ]
    )
    if as_text:
        sheet[
            "C2"
        ].number_format = "@"
    return _save_workbook(
        workbook
    )


def _normalized(
    raw: dict[str, object],
    mapping: dict[str, str],
) -> dict[str, object]:
    return normalize_raw_row(
        raw,
        mapping,
        default_lot_control_mode=
            "NONE",
        default_expiry_control_mode=
            "NONE",
    )


class _CaptureDb:
    def __init__(
        self,
    ) -> None:
        self.rows: list[
            dict[str, object]
        ] = []

    async def execute(
        self,
        statement,
    ) -> None:
        self.rows.extend(
            {
                getattr(key, "key", key): value
                for key, value in item.items()
            }
            for item in statement._multi_values[0]
        )


class Phase11RowFidelityTests(
    unittest.IsolatedAsyncioTestCase
):
    async def test_blank_xlsx_row_keeps_physical_row_number_through_staging(
        self,
    ) -> None:
        db = _CaptureDb()
        with open_source(
            "products.xlsx",
            _xlsx_with_blank_row(),
        ) as source:
            total = await insert_staged_rows(
                db,
                company_id=1,
                job_id=uuid4(),
                rows=source.rows,
                batch_size=1000,
                max_rows=50_000,
            )

        self.assertEqual(
            total,
            2,
        )
        self.assertEqual(
            [
                int(
                    row[
                        "row_number"
                    ]
                )
                for row in db.rows
            ],
            [
                2,
                4,
            ],
        )


class Phase11WorksheetTests(
    unittest.TestCase
):
    def test_external_workbook_with_multiple_visible_sheets_fails_closed(
        self,
    ) -> None:
        with self.assertRaisesRegex(
            ProductImportTerminalError,
            "multiple visible worksheets",
        ):
            with open_source(
                "external.xlsx",
                _external_multisheet_xlsx(),
            ):
                pass

    def test_official_template_marker_selects_canonical_products_sheet(
        self,
    ) -> None:
        payload = (
            build_product_import_template(
                locale="en"
            )
        )
        workbook = load_workbook(
            io.BytesIO(
                payload
            )
        )
        self.assertIn(
            WANASAH_TEMPLATE_META_SHEET,
            workbook.sheetnames,
        )
        meta = workbook[
            WANASAH_TEMPLATE_META_SHEET
        ]
        self.assertEqual(
            meta[
                "A1"
            ].value,
            WANASAH_TEMPLATE_MARKER,
        )
        canonical_sheet = str(
            meta[
                "A2"
            ].value
        )
        extra = workbook.create_sheet(
            "Visible notes"
        )
        extra[
            "A1"
        ] = "Not products"
        payload = _save_workbook(
            workbook
        )

        with open_source(
            "wanasah-template.xlsx",
            payload,
        ) as source:
            self.assertEqual(
                source.headers,
                [
                    EN_IMPORT_LOCALE.template.headers[
                        field
                    ]
                    for field
                    in CANONICAL_IMPORT_FIELDS
                ],
            )

        self.assertEqual(
            canonical_sheet,
            "Products",
        )


class Phase11FormulaTests(
    unittest.TestCase
):
    def test_formula_with_cached_value_uses_cached_value_for_non_identity_field(
        self,
    ) -> None:
        with open_source(
            "cached-formula.xlsx",
            _xlsx_formula_payload(
                cached_value="2"
            ),
        ) as source:
            row = next(
                source.rows
            )

        self.assertEqual(
            row.raw[
                "Unit Price"
            ],
            "2",
        )
        self.assertIn(
            "Unit Price",
            row.raw[
                SOURCE_CELL_META_KEY
            ],
        )
        normalized = _normalized(
            row.raw,
            {
                "name":
                    "Product",
                "unit_price":
                    "Unit Price",
            },
        )
        self.assertEqual(
            normalized[
                "unit_price"
            ],
            "2.000000",
        )

    def test_formula_without_cached_value_becomes_clear_row_error(
        self,
    ) -> None:
        with open_source(
            "uncached-formula.xlsx",
            _xlsx_formula_payload(
                cached_value=None
            ),
        ) as source:
            row = next(
                source.rows
            )

        with self.assertRaises(
            SimpleProductError
        ) as raised:
            _normalized(
                row.raw,
                {
                    "name":
                        "Product",
                    "unit_price":
                        "Unit Price",
                },
            )
        self.assertEqual(
            raised.exception.code,
            "IMPORT_FORMULA_VALUE_UNAVAILABLE",
        )


class Phase11BarcodeTests(
    unittest.TestCase
):
    def test_template_marks_barcode_columns_as_text_and_validates_text_input(
        self,
    ) -> None:
        payload = (
            build_product_import_template(
                locale="en"
            )
        )
        workbook = load_workbook(
            io.BytesIO(
                payload
            )
        )
        sheet = workbook[
            "Products"
        ]
        for field in (
            "unit_barcode",
            "package_barcode",
        ):
            column_index = (
                CANONICAL_IMPORT_FIELDS.index(
                    field
                )
                + 1
            )
            column_letter = chr(
                64
                + column_index
            )
            self.assertEqual(
                sheet.column_dimensions[
                    column_letter
                ].number_format,
                "@",
            )

        validations = list(
            sheet.data_validations.dataValidation
        )
        barcode_custom = [
            validation
            for validation in validations
            if validation.type
            == "custom"
            and "ISTEXT("
            in str(
                validation.formula1
            )
        ]
        self.assertEqual(
            len(
                barcode_custom
            ),
            2,
        )
        for validation in barcode_custom:
            self.assertIn(
                "ISTEXT(",
                str(
                    validation.formula1
                ),
            )
            self.assertIn(
                str(
                    MAX_TEMPLATE_ROWS
                    + 1
                ),
                str(
                    validation.sqref
                ),
            )
        workbook.close()

    def test_company_template_has_polished_layout_and_guided_entry_rules(
        self,
    ) -> None:
        payload = (
            build_product_import_template(
                locale="ar"
            )
        )
        workbook = load_workbook(
            io.BytesIO(
                payload
            )
        )
        sheet = workbook[
            "المنتجات"
        ]

        self.assertEqual(
            sheet["A1"].fill.fgColor.rgb,
            "00475569",
        )
        self.assertEqual(
            sheet.row_dimensions[
                1
            ].height,
            30,
        )
        self.assertEqual(
            sheet.column_dimensions[
                "E"
            ].width,
            14,
        )
        self.assertEqual(
            sheet.column_dimensions[
                "F"
            ].width,
            14,
        )
        self.assertEqual(
            sheet.column_dimensions[
                "I"
            ].width,
            16,
        )
        self.assertEqual(
            sheet.column_dimensions[
                "J"
            ].width,
            16,
        )
        self.assertEqual(
            sheet.column_dimensions[
                "D"
            ].number_format,
            "0",
        )
        self.assertEqual(
            sheet.column_dimensions[
                "E"
            ].number_format,
            "0.000",
        )
        self.assertEqual(
            sheet.column_dimensions[
                "F"
            ].number_format,
            "0.000",
        )

        validations = list(
            sheet.data_validations.dataValidation
        )
        package_lists = [
            validation
            for validation in validations
            if validation.type
            == "list"
            and "$B$1:$B$10"
            in str(
                validation.formula1
            )
        ]
        self.assertEqual(
            len(
                package_lists
            ),
            1,
        )
        self.assertFalse(
            package_lists[
                0
            ].allow_blank
        )
        self.assertIn(
            "C2:C50001",
            str(
                package_lists[
                    0
                ].sqref
            ),
        )

        unit_rules = [
            validation
            for validation in validations
            if validation.type
            == "custom"
            and "D2=INT(D2)"
            in str(
                validation.formula1
            )
        ]
        self.assertEqual(
            len(
                unit_rules
            ),
            1,
        )
        self.assertIn(
            "D2:D50001",
            str(
                unit_rules[
                    0
                ].sqref
            ),
        )

        package_barcode_rules = [
            validation
            for validation in validations
            if validation.type
            == "custom"
            and "H2"
            in str(
                validation.formula1
            )
            and "C2<>"
            in str(
                validation.formula1
            )
        ]
        self.assertEqual(
            len(
                package_barcode_rules
            ),
            1,
        )

        self.assertIsNotNone(
            sheet[
                "C1"
            ].comment
        )
        self.assertIsNotNone(
            sheet[
                "D1"
            ].comment
        )
        self.assertIsNotNone(
            sheet[
                "E1"
            ].comment
        )
        self.assertIsNotNone(
            sheet[
                "H1"
            ].comment
        )

        list_sheet = workbook[
            "_wanasah_lists"
        ]
        package_labels = [
            list_sheet.cell(
                row=row,
                column=2,
            ).value
            for row in range(
                1,
                11,
            )
        ]
        self.assertEqual(
            package_labels,
            [
                "بدون عبوة خارجية",
                "كرتونة",
                "صندوق",
                "باكيت",
                "كيس",
                "شوال",
                "صينية",
                "قفص",
                "حزمة",
                "طبلية",
            ],
        )
        workbook.close()

    def test_numeric_xlsx_barcode_is_rejected_as_unsafe(
        self,
    ) -> None:
        with open_source(
            "numeric-barcode.xlsx",
            _barcode_xlsx(
                1234567890123,
                as_text=False,
            ),
        ) as source:
            row = next(
                source.rows
            )

        with self.assertRaises(
            SimpleProductError
        ) as raised:
            _normalized(
                row.raw,
                {
                    "name":
                        "Product",
                    "unit_price":
                        "Unit Price",
                    "unit_barcode":
                        "Unit Barcode",
                },
            )
        self.assertEqual(
            raised.exception.code,
            "IMPORT_BARCODE_NUMERIC_UNSAFE",
        )

    def test_scientific_notation_barcode_text_is_rejected(
        self,
    ) -> None:
        raw = {
            "Product": "Tea",
            "Unit Price": "1.000",
            "Unit Barcode":
                "1.23456789E+12",
        }
        with self.assertRaises(
            SimpleProductError
        ) as raised:
            _normalized(
                raw,
                {
                    "name":
                        "Product",
                    "unit_price":
                        "Unit Price",
                    "unit_barcode":
                        "Unit Barcode",
                },
            )
        self.assertEqual(
            raised.exception.code,
            "IMPORT_BARCODE_SCIENTIFIC_NOTATION",
        )

    def test_leading_zero_text_barcode_is_preserved_exactly(
        self,
    ) -> None:
        with open_source(
            "text-barcode.xlsx",
            _barcode_xlsx(
                "001234567890",
                as_text=True,
            ),
        ) as source:
            row = next(
                source.rows
            )

        normalized = _normalized(
            row.raw,
            {
                "name":
                    "Product",
                "unit_price":
                    "Unit Price",
                "unit_barcode":
                    "Unit Barcode",
            },
        )
        self.assertEqual(
            normalized[
                "unit_barcode"
            ],
            "001234567890",
        )


    def test_formula_barcode_is_rejected_even_when_cached_value_exists(
        self,
    ) -> None:
        raw = {
            "Product": "Tea",
            "Unit Price": "1.000",
            "Unit Barcode": "001234",
            SOURCE_CELL_META_KEY: {
                "Unit Barcode": {
                    "formula": True,
                    "cached": True,
                }
            },
        }
        with self.assertRaises(
            SimpleProductError
        ) as raised:
            _normalized(
                raw,
                {
                    "name":
                        "Product",
                    "unit_price":
                        "Unit Price",
                    "unit_barcode":
                        "Unit Barcode",
                },
            )
        self.assertEqual(
            raised.exception.code,
            "IMPORT_BARCODE_FORMULA_NOT_ALLOWED",
        )


class Phase11FutureDateContractTests(
    unittest.TestCase
):
    def test_1900_epoch_rejects_phantom_serial_day_60(
        self,
    ) -> None:
        with self.assertRaisesRegex(
            ValueError,
            "serial day 60",
        ):
            normalize_spreadsheet_date(
                60,
                epoch="1900",
            )
        self.assertEqual(
            normalize_spreadsheet_date(
                61,
                epoch="1900",
            ),
            date(
                1900,
                3,
                1,
            ),
        )

    def test_1904_epoch_and_native_date_values_are_deterministic(
        self,
    ) -> None:
        self.assertEqual(
            normalize_spreadsheet_date(
                1,
                epoch="1904",
            ),
            date(
                1904,
                1,
                2,
            ),
        )
        self.assertEqual(
            normalize_spreadsheet_date(
                date(
                    2026,
                    9,
                    28,
                ),
                epoch="1900",
            ),
            date(
                2026,
                9,
                28,
            ),
        )
        self.assertEqual(
            normalize_spreadsheet_date(
                datetime(
                    2026,
                    9,
                    28,
                    0,
                    0,
                ),
                epoch="1900",
            ),
            date(
                2026,
                9,
                28,
            ),
        )

    def test_date_only_contract_accepts_iso_and_rejects_ambiguous_or_timezone_values(
        self,
    ) -> None:
        self.assertEqual(
            normalize_spreadsheet_date(
                "2026-09-28",
                epoch="1900",
            ),
            date(
                2026,
                9,
                28,
            ),
        )
        with self.assertRaises(
            ValueError
        ):
            normalize_spreadsheet_date(
                "09/28/2026",
                epoch="1900",
            )
        with self.assertRaises(
            ValueError
        ):
            normalize_spreadsheet_date(
                datetime(
                    2026,
                    9,
                    28,
                    tzinfo=timezone.utc,
                ),
                epoch="1900",
            )
        with self.assertRaises(
            ValueError
        ):
            normalize_spreadsheet_date(
                61.5,
                epoch="1900",
            )


if __name__ == "__main__":
    unittest.main()
