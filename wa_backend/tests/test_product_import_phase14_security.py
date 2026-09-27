from __future__ import annotations

import csv
from io import BytesIO, StringIO
from types import SimpleNamespace
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

from domains.simple_products.imports.application.correction_service import (
    _artifact_headers,
    _artifact_values,
    sanitize_spreadsheet_cell,
)
from domains.simple_products.imports.domain import (
    ProductImportTerminalError,
)
from domains.simple_products.imports.domain.errors import (
    import_error_field,
    public_error_summary,
    runtime_failure_summary,
    user_safe_row_error_message,
)
from domains.simple_products.imports.infrastructure.content_security import (
    validate_source_content,
)
from domains.simple_products.imports.infrastructure.parsers import (
    open_source,
)


def _zip_payload(
    *,
    macro: bool = False,
    bomb: bool = False,
) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "[Content_Types].xml",
            "<Types/>",
        )
        archive.writestr(
            "xl/workbook.xml",
            "<workbook/>",
        )
        if macro:
            archive.writestr(
                "xl/vbaProject.bin",
                b"macro",
            )
        if bomb:
            archive.writestr(
                "xl/sharedStrings.xml",
                b"A" * (2 * 1024 * 1024),
            )
    return output.getvalue()


class Phase14ErrorContractTests(unittest.TestCase):
    def test_runtime_summary_never_persists_technical_exception_text(self) -> None:
        secret = (
            "SELECT password FROM users; "
            "constraint uq_secret violated"
        )
        summary = runtime_failure_summary(
            message=secret,
            final_attempt=True,
            retryable=False,
            resume_status="IMPORTING",
            code="PRODUCT_IMPORT_SYSTEM_FAILURE",
            correlation_id="req-123",
        )

        self.assertNotIn(
            "technical",
            summary,
        )
        self.assertNotIn(
            secret,
            str(summary),
        )
        self.assertEqual(
            summary["correlation_id"],
            "req-123",
        )
        self.assertEqual(
            summary["code"],
            "PRODUCT_IMPORT_SYSTEM_FAILURE",
        )

    def test_public_summary_whitelists_legacy_storage(self) -> None:
        public = public_error_summary(
            {
                "code": "PRODUCT_IMPORT_SYSTEM_FAILURE",
                "message": "SQLSTATE 23505",
                "technical": "constraint secret",
                "stack": "traceback secret",
                "correlation_id": "req-456",
            }
        )
        self.assertEqual(
            set(public),
            {
                "code",
                "message",
                "correlation_id",
            },
        )
        self.assertNotIn(
            "SQLSTATE",
            str(public),
        )
        self.assertNotIn(
            "constraint secret",
            str(public),
        )

    def test_terminal_error_separates_internal_and_public_text(self) -> None:
        exc = ProductImportTerminalError(
            "SQLSTATE 23505 uq_private",
            code=
                "PRODUCT_IMPORT_SOURCE_TYPE_MISMATCH",
        )
        self.assertIn(
            "SQLSTATE",
            str(exc),
        )
        self.assertNotIn(
            "SQLSTATE",
            exc.user_message,
        )

    def test_row_contract_has_safe_code_message_and_actionable_field(self) -> None:
        self.assertEqual(
            import_error_field(
                "IMPORT_PACKAGE_BARCODE_WITHOUT_PACKAGE"
            ),
            "package_barcode",
        )
        self.assertEqual(
            user_safe_row_error_message(
                "IMPORT_BARCODE_CONFLICT"
            ),
            "The barcode is already in use.",
        )


class Phase14SpreadsheetInjectionTests(unittest.TestCase):
    def test_formula_like_cells_are_prefixed(self) -> None:
        for value in (
            "=1+1",
            "+cmd",
            "-2+3",
            "@SUM(A1:A2)",
            "  =hidden",
            "\t=hidden",
        ):
            with self.subTest(value=value):
                safe = sanitize_spreadsheet_cell(
                    value
                )
                self.assertTrue(
                    str(safe).startswith(
                        "'"
                    )
                )

        self.assertEqual(
            sanitize_spreadsheet_cell(
                "ordinary"
            ),
            "ordinary",
        )

    def test_correction_artifact_sanitizes_headers_source_and_error_cells(self) -> None:
        headers = _artifact_headers(
            [
                "=danger-header",
                "name",
            ]
        )
        self.assertIn(
            "'=danger-header",
            headers,
        )

        row = SimpleNamespace(
            row_identity=
                "11111111-1111-4111-8111-111111111111",
            row_number=2,
            error_code=
                "IMPORT_BARCODE_CONFLICT",
            error_message=
                "=HYPERLINK(\"http://evil\")",
            raw_data={
                "=danger-header":
                    "=WEBSERVICE(\"http://evil\")",
                "name":
                    "+SUM(1,1)",
            },
        )
        values = _artifact_values(
            row,
            [
                "=danger-header",
                "name",
            ],
        )

        self.assertEqual(
            values[3],
            "barcode",
        )
        self.assertEqual(
            values[4],
            "The barcode is already in use.",
        )
        self.assertTrue(
            str(values[5]).startswith("'"),
        )
        self.assertTrue(
            str(values[6]).startswith("'"),
        )
        self.assertNotIn(
            "HYPERLINK",
            str(values[4]),
        )

    def test_sanitized_xlsx_cell_is_literal_not_formula(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = sanitize_spreadsheet_cell(
            "=1+1"
        )
        output = BytesIO()
        workbook.save(output)
        workbook.close()

        loaded = load_workbook(
            BytesIO(
                output.getvalue()
            ),
            data_only=False,
        )
        try:
            cell = loaded.active["A1"]
            self.assertEqual(
                cell.value,
                "'=1+1",
            )
            self.assertNotEqual(
                cell.data_type,
                "f",
            )
        finally:
            loaded.close()


class Phase14FileSecurityTests(unittest.TestCase):
    def test_client_mime_is_irrelevant_and_csv_signature_is_textual(self) -> None:
        mime = validate_source_content(
            "products.csv",
            b"name,unit_price\nTea,1.0\n",
        )
        self.assertEqual(
            mime,
            "text/csv; charset=utf-8",
        )

    def test_binary_payload_disguised_as_csv_is_rejected(self) -> None:
        with self.assertRaises(
            ProductImportTerminalError
        ) as caught:
            validate_source_content(
                "products.csv",
                _zip_payload(),
            )
        self.assertEqual(
            caught.exception.code,
            "PRODUCT_IMPORT_SOURCE_TYPE_MISMATCH",
        )

    def test_nul_containing_csv_is_rejected(self) -> None:
        with self.assertRaises(
            ProductImportTerminalError
        ) as caught:
            open_source(
                "products.csv",
                b"name,price\nTea\x00,1\n",
            ).__enter__()
        self.assertEqual(
            caught.exception.code,
            "PRODUCT_IMPORT_SOURCE_TYPE_MISMATCH",
        )

    def test_malformed_xlsx_is_rejected_before_workbook_traversal(self) -> None:
        with self.assertRaises(
            ProductImportTerminalError
        ) as caught:
            validate_source_content(
                "products.xlsx",
                b"not-a-zip",
            )
        self.assertEqual(
            caught.exception.code,
            "PRODUCT_IMPORT_SOURCE_TYPE_MISMATCH",
        )

    def test_macro_payload_is_rejected(self) -> None:
        with self.assertRaises(
            ProductImportTerminalError
        ) as caught:
            validate_source_content(
                "products.xlsx",
                _zip_payload(
                    macro=True
                ),
            )
        self.assertEqual(
            caught.exception.code,
            "PRODUCT_IMPORT_XLSX_UNSAFE",
        )

    def test_high_compression_ratio_xlsx_is_rejected(self) -> None:
        with self.assertRaises(
            ProductImportTerminalError
        ) as caught:
            validate_source_content(
                "products.xlsx",
                _zip_payload(
                    bomb=True
                ),
            )
        self.assertEqual(
            caught.exception.code,
            "PRODUCT_IMPORT_XLSX_UNSAFE",
        )

    def test_csv_duplicate_headers_are_rejected_as_malformed(self) -> None:
        source = open_source(
            "products.csv",
            b"name,name\nTea,Tea\n",
        )
        with self.assertRaises(
            ProductImportTerminalError
        ):
            with source:
                pass


if __name__ == "__main__":
    unittest.main()
