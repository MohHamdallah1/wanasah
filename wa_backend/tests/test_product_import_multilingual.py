from __future__ import annotations

from dataclasses import replace
import io
import unittest

from openpyxl import load_workbook

from domains.simple_products.imports.domain import (
    AR_IMPORT_LOCALE,
    EN_IMPORT_LOCALE,
    build_import_alias_registry,
    build_import_locale_registry,
    resolve_import_locale_pack,
    suggest_import_mapping,
)
from domains.simple_products.imports.domain.normalization import (
    normalize_raw_row,
)
from domains.simple_products.imports.infrastructure.template import (
    build_product_import_template,
)


class ProductImportLocaleResolutionTests(
    unittest.TestCase
):
    def test_ui_locale_tags_resolve_by_exact_then_primary_language(
        self,
    ) -> None:
        self.assertIs(
            resolve_import_locale_pack(
                "ar-JO"
            ),
            AR_IMPORT_LOCALE,
        )
        self.assertIs(
            resolve_import_locale_pack(
                "en-US"
            ),
            EN_IMPORT_LOCALE,
        )
        self.assertIs(
            resolve_import_locale_pack(
                "AR_jo"
            ),
            AR_IMPORT_LOCALE,
        )

    def test_unknown_ui_locale_falls_back_without_binding_parser_language(
        self,
    ) -> None:
        self.assertIs(
            resolve_import_locale_pack(
                "fr-FR"
            ),
            EN_IMPORT_LOCALE,
        )

    def test_adding_future_locale_pack_requires_registry_data_not_parser_changes(
        self,
    ) -> None:
        french = replace(
            EN_IMPORT_LOCALE,
            locale="fr",
        )
        registry = (
            build_import_locale_registry(
                (
                    EN_IMPORT_LOCALE,
                    AR_IMPORT_LOCALE,
                    french,
                )
            )
        )
        self.assertIs(
            resolve_import_locale_pack(
                "fr-FR",
                registry=registry,
            ),
            french,
        )


class ProductImportMixedLanguageTests(
    unittest.TestCase
):
    def test_arabic_headers_accept_english_controlled_values(
        self,
    ) -> None:
        raw = {
            "اسم المنتج":
                "شاي أسود",
            "نوع العبوة":
                "carton",
            "عدد الوحدات في العبوة":
                "12",
            "سعر الوحدة":
                "1.250",
            "تتبع الدفعة":
                "required",
            "تتبع الصلاحية":
                "none",
        }
        mapping = suggest_import_mapping(
            list(
                raw
            )
        )
        normalized = normalize_raw_row(
            raw,
            mapping,
            default_lot_control_mode=
                "OPTIONAL",
            default_expiry_control_mode=
                "OPTIONAL",
        )

        self.assertEqual(
            mapping[
                "name"
            ],
            "اسم المنتج",
        )
        self.assertEqual(
            normalized[
                "package_uom_code"
            ],
            "CARTON",
        )
        self.assertEqual(
            normalized[
                "lot_control_mode"
            ],
            "REQUIRED",
        )
        self.assertEqual(
            normalized[
                "expiry_control_mode"
            ],
            "NONE",
        )

    def test_english_headers_accept_arabic_controlled_values(
        self,
    ) -> None:
        raw = {
            "Product Name":
                "Coffee قهوة",
            "Package Type":
                "كرتونة",
            "Units per Package":
                "24",
            "Unit Price":
                "2.500",
            "Lot Tracking":
                "إلزامي",
            "Expiry Tracking":
                "بدون",
        }
        mapping = suggest_import_mapping(
            list(
                raw
            )
        )
        normalized = normalize_raw_row(
            raw,
            mapping,
            default_lot_control_mode=
                "OPTIONAL",
            default_expiry_control_mode=
                "OPTIONAL",
        )

        self.assertEqual(
            mapping[
                "package_uom"
            ],
            "Package Type",
        )
        self.assertEqual(
            normalized[
                "name"
            ],
            "Coffee قهوة",
        )
        self.assertEqual(
            normalized[
                "package_uom_code"
            ],
            "CARTON",
        )
        self.assertEqual(
            normalized[
                "lot_control_mode"
            ],
            "REQUIRED",
        )
        self.assertEqual(
            normalized[
                "expiry_control_mode"
            ],
            "NONE",
        )

    def test_mixed_language_headers_share_one_global_alias_registry(
        self,
    ) -> None:
        headers = [
            "Product Name",
            "العائلة",
            "Package Type",
            "عدد الوحدات في العبوة",
            "Unit Price",
            "باركود الوحدة",
        ]
        mapping = suggest_import_mapping(
            headers
        )
        self.assertEqual(
            mapping,
            {
                "name":
                    "Product Name",
                "family":
                    "العائلة",
                "package_uom":
                    "Package Type",
                "units_per_package":
                    "عدد الوحدات في العبوة",
                "unit_price":
                    "Unit Price",
                "unit_barcode":
                    "باركود الوحدة",
            },
        )

    def test_alias_registry_contains_all_registered_languages_simultaneously(
        self,
    ) -> None:
        registry = (
            build_import_alias_registry(
                (
                    EN_IMPORT_LOCALE,
                    AR_IMPORT_LOCALE,
                )
            )
        )
        self.assertEqual(
            registry.header_to_field[
                "product name"
            ],
            "name",
        )
        self.assertEqual(
            registry.header_to_field[
                "اسم المنتج"
            ],
            "name",
        )
        self.assertEqual(
            registry.package_to_code[
                "carton"
            ],
            "CARTON",
        )
        self.assertEqual(
            registry.package_to_code[
                "كرتونة"
            ],
            "CARTON",
        )


class ProductImportTemplateLanguageTests(
    unittest.TestCase
):
    def test_regional_arabic_locale_builds_arabic_rtl_template(
        self,
    ) -> None:
        workbook = load_workbook(
            io.BytesIO(
                build_product_import_template(
                    locale="ar-JO"
                )
            )
        )
        try:
            sheet = workbook[
                "المنتجات"
            ]
            self.assertTrue(
                sheet.sheet_view.rightToLeft
            )
            self.assertEqual(
                sheet[
                    "A1"
                ].value,
                "اسم المنتج",
            )
        finally:
            workbook.close()

    def test_regional_english_locale_builds_english_ltr_template(
        self,
    ) -> None:
        workbook = load_workbook(
            io.BytesIO(
                build_product_import_template(
                    locale="en-US"
                )
            )
        )
        try:
            sheet = workbook[
                "Products"
            ]
            self.assertFalse(
                bool(
                    sheet.sheet_view.rightToLeft
                )
            )
            self.assertEqual(
                sheet[
                    "A1"
                ].value,
                "Name",
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
