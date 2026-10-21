"""Safe additive V1 import cost: original 10 columns/engine remain compatible."""
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook, Workbook

from domains.simple_products.imports.domain.localization import CANONICAL_IMPORT_FIELDS
from domains.simple_products.imports.domain.mapping import suggest_mapping, validate_mapping
from domains.simple_products.imports.domain.normalization import (
    normalize_raw_row, optional_purchase_unit_cost,
)
from domains.simple_products.imports.infrastructure.template import build_product_import_template
from domains.simple_products.imports.infrastructure.parsers import open_source
from domains.simple_products.imports.application.purchase_reference_import import has_import_cost
from domains.simple_products.service import SimpleProductError


LEGACY = (
    "name", "family", "package_uom", "units_per_package",
    "package_price", "unit_price", "unit_barcode", "package_barcode",
    "lot_control_mode", "expiry_control_mode",
)


@pytest.mark.parametrize("locale", ["ar", "en"])
def test_new_template_preserves_original_ten_columns_and_hidden_lists(locale):
    template = load_workbook(BytesIO(build_product_import_template(locale=locale)))
    sheet = template.active
    assert tuple(CANONICAL_IMPORT_FIELDS[:10]) == LEGACY
    assert CANONICAL_IMPORT_FIELDS[10] == "purchase_unit_cost"
    assert sheet.max_column == 11
    assert sheet.auto_filter.ref == "A1:K50001"
    assert template["_wanasah_meta"].sheet_state == "veryHidden"
    assert template["_wanasah_lists"].sheet_state == "hidden"
    assert template["_wanasah_meta"]["A1"].value == "WANASAH_PRODUCT_IMPORT_V1"
    assert [sheet.cell(1, i).value for i in range(1, 11)] == [
        __import__("domains.simple_products.imports.domain.localization",
                   fromlist=["resolve_import_locale_pack"]).resolve_import_locale_pack(locale)
        .template.headers[k] for k in LEGACY
    ]
    assert "تكلفة" in sheet["K1"].value if locale == "ar" else "Purchase Cost" in sheet["K1"].value
    assert any(str(r.sqref) == "K2:K50001" for r in sheet.data_validations.dataValidation)
    template.close()


def test_old_ten_column_workbooks_remain_acceptable_and_cost_absent():
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Products"
    sheet.append(["Name", "Family", "Outer Package Type", "Base Units per Outer Package",
                  "Outer Package Price", "Unit Price", "Unit Barcode", "Outer Package Barcode",
                  "Lot Tracking", "Expiry Tracking"])
    sheet.append(["Legacy import", "", "NONE", "", "", "2.50", "", "", "NONE", "NONE"])
    f = BytesIO()
    wb.save(f)
    wb.close()
    with open_source("old-10-columns.xlsx", f.getvalue()) as source:
        assert len(source.headers) == 10
        mapping = validate_mapping(source.headers, suggest_mapping(source.headers))
        assert "purchase_unit_cost" not in mapping
        data = normalize_raw_row(next(iter(source.rows)).raw, mapping,
                                 default_lot_control_mode="NONE",
                                 default_expiry_control_mode="NONE")
    assert data["purchase_unit_cost"] is None
    assert data["unit_price"] == "2.500000"
    assert has_import_cost([SimpleNamespace(normalized_data=data)]) is False


@pytest.mark.parametrize("cost,expected", [
    ("0", "0.000000"), ("1.2", "1.200000"), ("12.123456", "12.123456"),
    ("99999999999999.999999", "99999999999999.999999"),
])
def test_optional_purchase_cost_exact_decimal(cost, expected):
    assert optional_purchase_unit_cost(cost) == expected


@pytest.mark.parametrize("cost", [
    "-1", "NaN", "Infinity", "1e10", "1.1234567", "100000000000000",
    "4,2", "1+2", True, "1.2.3",
])
def test_invalid_purchase_cost_is_row_invalid_not_rounded(cost):
    with pytest.raises(SimpleProductError, match="purchase|Purchase"):
        optional_purchase_unit_cost(cost)


def test_import_cost_does_not_infer_from_selling_price():
    data = {"unit_price": "250.000000", "purchase_unit_cost": None}
    assert has_import_cost([SimpleNamespace(normalized_data=data)]) is False
    data["purchase_unit_cost"] = "0.000000"
    assert has_import_cost([SimpleNamespace(normalized_data=data)]) is True
