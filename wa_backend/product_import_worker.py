"""Worker-side parse, staging, validation and import orchestration."""
from __future__ import annotations

import asyncio
import csv
import io
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any
from uuid import UUID, uuid5

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import delete, func, insert, select, text

from context import tenant_context
from database import AsyncSessionLocal
from domains.product_tracking import (
    ProductTrackingError,
    normalize_tracking_mode,
)
from domains.simple_products.service import (
    SimpleProductError,
    SimpleProductSpec,
    create_products_and_prices,
    normalize_barcode,
    normalize_package_code,
    resolve_price_pair,
)
from inventory_access import InventoryAccess
from models import (
    Driver,
    ProductBarcode,
    ProductImportJob,
    ProductImportRow,
)

MAX_IMPORT_ROWS = 50_000
MAX_IMPORT_COLUMNS = 100
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 500
MAX_XLSX_COMPRESSION_RATIO = 200
STAGE_BATCH = 1_000
IMPORT_BATCH = 100


class ProductImportTerminalError(RuntimeError):
    """Deterministic import failure that must not be retried automatically."""


_CANONICAL_FIELDS = (
    "name",
    "family",
    "package_uom",
    "units_per_package",
    "package_price",
    "unit_price",
    "unit_barcode",
    "package_barcode",
    "lot_control_mode",
    "expiry_control_mode",
)

_ALIASES = {
    "name": {
        "name",
        "product name",
        "product",
        "item",
        "item name",
        "اسم المنتج",
        "المنتج",
        "اسم الصنف",
        "الصنف",
    },
    "family": {
        "family",
        "product family",
        "group",
        "العائلة",
        "عائلة",
        "عائلة المنتج",
    },
    "package_uom": {
        "package",
        "package type",
        "outer package",
        "package uom",
        "نوع العبوة",
        "العبوة",
        "وحدة العبوة",
    },
    "units_per_package": {
        "units per package",
        "units/package",
        "pieces per package",
        "units per carton",
        "units/carton",
        "pieces per carton",
        "pcs per carton",
        "pack per carton",
        "packs per carton",
        "عدد الوحدات في العبوة",
        "عدد الحبات في العبوة",
        "عدد الحبات في الكرتونة",
        "الحبات بالكرتونة",
        "حبة بالكرتونة",
        "عدد القطع في الكرتونة",
    },
    "package_price": {
        "package price",
        "outer price",
        "carton price",
        "case price",
        "سعر العبوة",
        "سعر الكرتونة",
        "سعر كرتونة",
    },
    "unit_price": {
        "unit price",
        "piece price",
        "each price",
        "سعر الوحدة",
        "سعر الحبة",
        "سعر القطعة",
    },
    "unit_barcode": {
        "unit barcode",
        "piece barcode",
        "each barcode",
        "باركود الوحدة",
        "باركود الحبة",
        "باركود القطعة",
        "باركود",
    },
    "package_barcode": {
        "package barcode",
        "outer barcode",
        "carton barcode",
        "case barcode",
        "باركود العبوة",
        "باركود الكرتونة",
    },
    "lot_control_mode": {
        "lot control",
        "lot tracking",
        "batch tracking",
        "lot control mode",
        "batch control mode",
        "تتبع الدفعة",
        "تتبع الدفعات",
        "تتبع التشغيلة",
        "نمط تتبع الدفعة",
    },
    "expiry_control_mode": {
        "expiry control",
        "expiry tracking",
        "expiration tracking",
        "expiry control mode",
        "expiration control mode",
        "تتبع الصلاحية",
        "تتبع تاريخ الصلاحية",
        "نمط تتبع الصلاحية",
    },
}

_TRACKING_VALUE_ALIASES = {
    "none": "NONE",
    "no": "NONE",
    "without": "NONE",
    "بدون": "NONE",
    "لا": "NONE",
    "optional": "OPTIONAL",
    "اختياري": "OPTIONAL",
    "required": "REQUIRED",
    "mandatory": "REQUIRED",
    "إلزامي": "REQUIRED",
    "الزامي": "REQUIRED",
}


def _tracking_import_mode(
    raw_value: Any,
    *,
    fallback: str,
    field_name: str,
) -> str:
    if raw_value is None or not str(raw_value).strip():
        return normalize_tracking_mode(
            fallback,
            field_name=field_name,
        )

    normalized_value = _normalize_header(raw_value)
    canonical = _TRACKING_VALUE_ALIASES.get(
        normalized_value,
        str(raw_value).strip(),
    )
    return normalize_tracking_mode(
        canonical,
        field_name=field_name,
    )


_PACKAGE_VALUE_ALIASES = {
    "carton": "CARTON",
    "كرتونة": "CARTON",
    "case": "CASE",
    "box": "CASE",
    "صندوق": "CASE",
    "pack": "PACK",
    "packet": "PACK",
    "باكيت": "PACK",
    "bag": "BAG",
    "كيس": "BAG",
    "sack": "SACK",
    "شوال": "SACK",
    "tray": "TRAY",
    "صينية": "TRAY",
    "crate": "CRATE",
    "قفص": "CRATE",
    "bundle": "BUNDLE",
    "حزمة": "BUNDLE",
    "pallet": "PALLET",
    "طبلية": "PALLET",
    "none": "NONE",
    "no package": "NONE",
    "unit only": "NONE",
    "بدون": "NONE",
    "بدون عبوة": "NONE",
    "لا يوجد": "NONE",
}


def _normalize_header(value: Any) -> str:
    normalized = (
        str(value or "")
        .strip()
        .lower()
        .replace("_", " ")
        .replace("-", " ")
    )
    return re.sub(r"\s+", " ", normalized)


_NORMALIZED_ALIASES = {
    key: {
        _normalize_header(value)
        for value in aliases
    }
    for key, aliases in _ALIASES.items()
}


def _json_cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    clean = str(value).strip()
    return clean or None


def _header_layout(
    values: list[Any] | tuple[Any, ...],
) -> list[tuple[int, str]]:
    layout = [
        (
            index,
            str(value or "").strip(),
        )
        for index, value in enumerate(values)
        if str(value or "").strip()
    ]
    if not layout:
        raise ProductImportTerminalError(
            "The header row is empty."
        )
    if len(layout) > MAX_IMPORT_COLUMNS:
        raise ProductImportTerminalError(
            f"The file has more than {MAX_IMPORT_COLUMNS} columns."
        )

    normalized = [
        _normalize_header(header)
        for _, header in layout
    ]
    if len(normalized) != len(set(normalized)):
        raise ProductImportTerminalError(
            "The header row contains duplicate columns after normalization."
        )
    return layout


def _row_from_layout(
    values: list[Any] | tuple[Any, ...],
    layout: list[tuple[int, str]],
) -> dict[str, Any]:
    return {
        header: _json_cell(
            values[index]
            if index < len(values)
            else None
        )
        for index, header in layout
    }


def _parse_csv(
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    decoded = None
    last_error = None
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1256",
    ):
        try:
            decoded = payload.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc

    if decoded is None:
        raise ProductImportTerminalError(
            "CSV encoding is not supported."
        ) from last_error

    try:
        dialect = csv.Sniffer().sniff(
            decoded[:8192],
            delimiters=",;\t|",
        )
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(
        io.StringIO(decoded),
        dialect=dialect,
    )
    first = next(reader, None)
    if first is None:
        raise ProductImportTerminalError(
            "The file is empty."
        )

    layout = _header_layout(first)
    headers = [
        header
        for _, header in layout
    ]
    rows: list[dict[str, Any]] = []
    for values in reader:
        raw = _row_from_layout(
            values,
            layout,
        )
        if not any(
            value is not None
            for value in raw.values()
        ):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ProductImportTerminalError(
                f"The import exceeds the {MAX_IMPORT_ROWS:,}-row safety limit."
            )
        rows.append(raw)

    return headers, rows


def _validate_xlsx_archive(
    payload: bytes,
) -> None:
    try:
        with zipfile.ZipFile(
            io.BytesIO(payload)
        ) as archive:
            infos = archive.infolist()
            if not infos:
                raise ProductImportTerminalError(
                    "The Excel file is empty or invalid."
                )
            if (
                len(infos)
                > MAX_XLSX_ARCHIVE_ENTRIES
            ):
                raise ProductImportTerminalError(
                    "The Excel archive has an abnormal number of internal entries."
                )

            total = 0
            for info in infos:
                total += int(info.file_size)
                if (
                    total
                    > MAX_XLSX_UNCOMPRESSED_BYTES
                ):
                    raise ProductImportTerminalError(
                        "The uncompressed Excel content exceeds the safety limit."
                    )
                if (
                    info.compress_size > 0
                    and info.file_size > 1_000_000
                    and (
                        info.file_size
                        / info.compress_size
                    )
                    > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise ProductImportTerminalError(
                        "The Excel file was rejected because of an abnormal compression ratio."
                    )
                if info.filename.lower().endswith(
                    "vbaproject.bin"
                ):
                    raise ProductImportTerminalError(
                        "Macro-enabled Excel files are not supported."
                    )
    except zipfile.BadZipFile as exc:
        raise ProductImportTerminalError(
            "The Excel file is invalid or corrupted."
        ) from exc


def _parse_xlsx(
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    _validate_xlsx_archive(payload)
    try:
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise ProductImportTerminalError(
            "The Excel file could not be opened."
        ) from exc

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(
            values_only=True
        )
        first = next(iterator, None)
        if first is None:
            raise ProductImportTerminalError(
                "The file is empty."
            )

        layout = _header_layout(first)
        headers = [
            header
            for _, header in layout
        ]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            raw = _row_from_layout(
                values,
                layout,
            )
            if not any(
                value is not None
                for value in raw.values()
            ):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ProductImportTerminalError(
                    f"The import exceeds the {MAX_IMPORT_ROWS:,}-row safety limit."
                )
            rows.append(raw)

        return headers, rows
    finally:
        workbook.close()


def parse_source(
    file_name: str,
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = (
        file_name.lower().rsplit(".", 1)[-1]
        if "." in file_name
        else ""
    )
    if suffix == "csv":
        return _parse_csv(payload)
    if suffix == "xlsx":
        return _parse_xlsx(payload)
    raise ProductImportTerminalError(
        "Only CSV and XLSX files are supported."
    )


def suggest_mapping(
    headers: list[str],
) -> dict[str, str]:
    normalized = {
        header: _normalize_header(header)
        for header in headers
    }
    result: dict[str, str] = {}
    for canonical in _CANONICAL_FIELDS:
        matches = [
            header
            for header, value in normalized.items()
            if value
            in _NORMALIZED_ALIASES[
                canonical
            ]
        ]
        if len(matches) == 1:
            result[canonical] = matches[0]
    return result


def mapping_complete(
    mapping: dict[str, str],
) -> bool:
    return bool(
        mapping.get("name")
        and (
            mapping.get("package_price")
            or mapping.get("unit_price")
        )
    )


def validate_mapping(
    headers: list[str],
    mapping: dict[str, str],
) -> dict[str, str]:
    allowed = set(headers)
    cleaned: dict[str, str] = {}
    used: set[str] = set()

    for canonical, header in mapping.items():
        if canonical not in _CANONICAL_FIELDS:
            raise ValueError(
                f"Unknown mapping field: {canonical}"
            )
        if not header:
            continue
        if header not in allowed:
            raise ValueError(
                f"Source column does not exist: {header}"
            )
        if header in used:
            raise ValueError(
                "One source column cannot map to more than one canonical field."
            )
        cleaned[canonical] = header
        used.add(header)

    return cleaned


def _package_code(
    raw_package: Any,
    *,
    has_package_column: bool,
    has_units_column: bool,
) -> str | None:
    if raw_package is not None and str(
        raw_package
    ).strip():
        normalized = _normalize_header(
            raw_package
        )
        canonical = _PACKAGE_VALUE_ALIASES.get(
            normalized,
            str(raw_package).strip().upper(),
        )
        return normalize_package_code(canonical)

    # Backward-compatible old carton files normally carry a carton quantity
    # column but no explicit package-type column.
    if has_units_column:
        return "CARTON"

    # No package column and no package-size column is explicit unit-only input.
    if not has_package_column:
        return None

    return None


def normalize_raw_row(
    raw: dict[str, Any],
    mapping: dict[str, str],
    *,
    default_lot_control_mode: str,
    default_expiry_control_mode: str,
) -> dict[str, Any]:
    def value(field: str):
        header = mapping.get(field)
        return (
            raw.get(header)
            if header
            else None
        )

    name = str(
        value("name") or ""
    ).strip()
    if not name:
        raise SimpleProductError(
            "IMPORT_NAME_REQUIRED",
            "Product name is required.",
            status_code=422,
        )
    if len(name) > 200:
        raise SimpleProductError(
            "IMPORT_NAME_TOO_LONG",
            "Product name exceeds the supported length.",
            status_code=422,
        )

    has_package_column = bool(
        mapping.get("package_uom")
    )
    has_units_column = bool(
        mapping.get("units_per_package")
    )
    package_code = _package_code(
        value("package_uom"),
        has_package_column=has_package_column,
        has_units_column=has_units_column,
    )

    if package_code is None:
        units = 1
    else:
        units_raw = str(
            value("units_per_package") or ""
        ).strip()
        if not units_raw:
            raise SimpleProductError(
                "IMPORT_PACKAGING_REQUIRED",
                "units_per_package is required for an outer package.",
                status_code=422,
            )
        try:
            units_decimal = Decimal(
                units_raw
            )
            units = int(units_decimal)
        except Exception as exc:
            raise SimpleProductError(
                "IMPORT_PACKAGING_INVALID",
                "units_per_package is invalid.",
                status_code=422,
            ) from exc

        if (
            units_decimal
            != Decimal(units)
            or units < 2
            or units > 1_000_000
        ):
            raise SimpleProductError(
                "IMPORT_PACKAGING_INVALID",
                "units_per_package must be an integer between 2 and 1,000,000.",
                status_code=422,
            )

    prices = resolve_price_pair(
        units_per_package=units,
        package_uom_code=package_code,
        package_price=value(
            "package_price"
        ),
        unit_price=value("unit_price"),
    )

    family = str(
        value("family") or ""
    ).strip() or None
    unit_barcode = normalize_barcode(
        value("unit_barcode"),
        "unit_barcode",
    )
    package_barcode = normalize_barcode(
        value("package_barcode"),
        "package_barcode",
    )
    if (
        package_code is None
        and package_barcode is not None
    ):
        raise SimpleProductError(
            "IMPORT_PACKAGE_BARCODE_WITHOUT_PACKAGE",
            "package_barcode cannot be supplied without an outer package.",
            status_code=422,
        )

    lot_control_mode = _tracking_import_mode(
        value("lot_control_mode"),
        fallback=default_lot_control_mode,
        field_name="lot_control_mode",
    )
    expiry_control_mode = _tracking_import_mode(
        value("expiry_control_mode"),
        fallback=default_expiry_control_mode,
        field_name="expiry_control_mode",
    )

    return {
        "name": name,
        "family_name": family,
        "package_uom_code": package_code,
        "units_per_package": units,
        "package_price": (
            format(
                prices.package_price,
                "f",
            )
            if prices.package_price
            is not None
            else None
        ),
        "unit_price": format(
            prices.unit_price,
            "f",
        ),
        "unit_barcode": unit_barcode,
        "package_barcode": package_barcode,
        "lot_control_mode": lot_control_mode,
        "expiry_control_mode": expiry_control_mode,
    }


async def _tenant_session(
    company_id: int,
):
    token = tenant_context.set(
        int(company_id)
    )
    db = AsyncSessionLocal()
    try:
        await db.execute(
            text(
                "SELECT set_config("
                "'app.current_tenant', :c, false)"
            ),
            {
                "c": str(
                    int(company_id)
                )
            },
        )
        return token, db
    except Exception:
        await db.close()
        tenant_context.reset(token)
        raise


async def _close(
    token,
    db,
) -> None:
    try:
        await db.close()
    finally:
        tenant_context.reset(token)


async def _set_status(
    company_id: int,
    job_id: UUID,
    status: str,
    **values: Any,
) -> None:
    token, db = await _tenant_session(
        company_id
    )
    try:
        job = await db.scalar(
            select(ProductImportJob)
            .where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
            .with_for_update()
        )
        if job is None:
            raise ValueError(
                "Product import job not found."
            )

        job.status = status
        for key, value in values.items():
            setattr(
                job,
                key,
                value,
            )
        job.version += 1
        job.updated_at = (
            datetime.now(
                timezone.utc
            ).replace(
                tzinfo=None
            )
        )
        await db.commit()
    except Exception:
        await db.rollback()
        raise
    finally:
        await _close(
            token,
            db,
        )


async def _load_job(
    company_id: int,
    job_id: UUID,
):
    token, db = await _tenant_session(
        company_id
    )
    try:
        job = await db.scalar(
            select(ProductImportJob).where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
        )
        if job is None:
            raise ValueError(
                "Product import job not found."
            )

        return (
            bytes(job.source_payload)
            if job.source_payload
            is not None
            else None,
            str(job.file_name),
            str(job.status),
            list(
                job.detected_headers
                or []
            ),
            dict(
                job.column_mapping
                or {}
            ),
        )
    finally:
        await _close(
            token,
            db,
        )


async def _stage_source(
    company_id: int,
    job_id: UUID,
    headers: list[str],
    rows: list[dict[str, Any]],
    suggestions: dict[str, str],
) -> None:
    if len(rows) > MAX_IMPORT_ROWS:
        raise ProductImportTerminalError(
            f"The import exceeds the {MAX_IMPORT_ROWS:,}-row safety limit."
        )

    token, db = await _tenant_session(
        company_id
    )
    try:
        await db.execute(
            delete(
                ProductImportRow
            ).where(
                ProductImportRow.company_id
                == int(company_id),
                ProductImportRow.job_id
                == job_id,
            )
        )

        for offset in range(
            0,
            len(rows),
            STAGE_BATCH,
        ):
            chunk = rows[
                offset : offset + STAGE_BATCH
            ]
            await db.execute(
                insert(ProductImportRow),
                [
                    {
                        "company_id": int(
                            company_id
                        ),
                        "job_id": job_id,
                        "row_number": (
                            offset
                            + index
                            + 2
                        ),
                        "raw_data": raw,
                        "normalized_data": {},
                        "status": "STAGED",
                        "version": 1,
                    }
                    for index, raw
                    in enumerate(chunk)
                ],
            )

        job = await db.scalar(
            select(ProductImportJob)
            .where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
            .with_for_update()
        )
        if job is None:
            raise ValueError(
                "Product import job not found."
            )

        job.detected_headers = headers
        job.suggested_mapping = suggestions
        if not job.column_mapping:
            job.column_mapping = (
                suggestions
            )
        job.total_rows = len(rows)
        job.processed_rows = 0
        job.valid_rows = 0
        job.failed_rows = 0
        job.source_payload = None
        job.status = (
            "VALIDATING"
            if mapping_complete(
                dict(
                    job.column_mapping
                    or {}
                )
            )
            else "NEEDS_MAPPING"
        )
        job.version += 1
        job.updated_at = (
            datetime.now(
                timezone.utc
            ).replace(
                tzinfo=None
            )
        )
        await db.commit()
    except Exception:
        await db.rollback()
        raise
    finally:
        await _close(
            token,
            db,
        )


async def _validate_rows(
    company_id: int,
    job_id: UUID,
) -> bool:
    token, db = await _tenant_session(
        company_id
    )
    try:
        job = await db.scalar(
            select(ProductImportJob).where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
        )
        if job is None:
            raise ValueError(
                "Product import job not found."
            )

        try:
            mapping = validate_mapping(
                list(
                    job.detected_headers
                    or []
                ),
                dict(
                    job.column_mapping
                    or {}
                ),
            )
        except ValueError as exc:
            raise ProductImportTerminalError(
                str(exc)
            ) from exc

        if not mapping_complete(mapping):
            await db.rollback()
            await _set_status(
                company_id,
                job_id,
                "NEEDS_MAPPING",
            )
            return False

        rows = list(
            (
                await db.scalars(
                    select(ProductImportRow)
                    .where(
                        ProductImportRow.company_id
                        == int(company_id),
                        ProductImportRow.job_id
                        == job_id,
                    )
                    .order_by(
                        ProductImportRow.row_number.asc()
                    )
                )
            ).all()
        )

        normalized: dict[
            int,
            dict[str, Any],
        ] = {}
        errors: dict[
            int,
            tuple[str, str],
        ] = {}
        barcode_rows: dict[
            str,
            list[int],
        ] = {}

        for row in rows:
            try:
                data = normalize_raw_row(
                    dict(
                        row.raw_data
                        or {}
                    ),
                    mapping,
                    default_lot_control_mode=str(
                        job.default_lot_control_mode
                    ),
                    default_expiry_control_mode=str(
                        job.default_expiry_control_mode
                    ),
                )
                row_id = int(row.id)
                normalized[
                    row_id
                ] = data

                # Same unit/package barcode in ONE row is intentional shared
                # identity. It only becomes a duplicate if it spans rows.
                row_barcodes = {
                    str(barcode)
                    for barcode in (
                        data.get(
                            "unit_barcode"
                        ),
                        data.get(
                            "package_barcode"
                        ),
                    )
                    if barcode
                }
                for barcode in row_barcodes:
                    barcode_rows.setdefault(
                        barcode,
                        [],
                    ).append(
                        row_id
                    )
            except ProductTrackingError as exc:
                errors[int(row.id)] = (
                    exc.code,
                    exc.message,
                )
            except SimpleProductError as exc:
                errors[int(row.id)] = (
                    exc.code,
                    exc.message,
                )
            except Exception as exc:
                errors[int(row.id)] = (
                    "IMPORT_ROW_INVALID",
                    str(exc),
                )

        for barcode, row_ids in (
            barcode_rows.items()
        ):
            if len(row_ids) > 1:
                for row_id in row_ids:
                    errors[row_id] = (
                        "IMPORT_BARCODE_DUPLICATE",
                        f"Barcode {barcode} appears on more than one import row.",
                    )

        candidates = [
            barcode
            for barcode, row_ids
            in barcode_rows.items()
            if len(row_ids) == 1
        ]
        if candidates:
            existing = set(
                (
                    await db.scalars(
                        select(
                            ProductBarcode.barcode
                        ).where(
                            ProductBarcode.company_id
                            == int(company_id),
                            ProductBarcode.barcode.in_(
                                candidates
                            ),
                            ProductBarcode.is_active.is_(
                                True
                            ),
                        )
                    )
                ).all()
            )
            for barcode in existing:
                for row_id in barcode_rows.get(
                    str(barcode),
                    [],
                ):
                    errors[
                        row_id
                    ] = (
                        "IMPORT_BARCODE_CONFLICT",
                        f"Barcode {barcode} is already active.",
                    )

        for row in rows:
            row_id = int(row.id)
            row.normalized_data = (
                normalized.get(
                    row_id,
                    {},
                )
            )
            if row_id in errors:
                row.status = "FAILED"
                (
                    row.error_code,
                    row.error_message,
                ) = errors[row_id]
            else:
                row.status = "VALID"
                row.error_code = None
                row.error_message = None
            row.version += 1

        valid_count = sum(
            1
            for row in rows
            if row.status == "VALID"
        )
        failed_count = (
            len(rows)
            - valid_count
        )

        locked = await db.scalar(
            select(ProductImportJob)
            .where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
            .with_for_update()
        )
        assert locked is not None
        locked.column_mapping = mapping
        locked.valid_rows = valid_count
        locked.failed_rows = failed_count
        locked.processed_rows = 0
        locked.status = (
            "VALIDATION_FAILED"
            if failed_count
            else "IMPORTING"
        )
        locked.error_summary = (
            {
                "code": "PRODUCT_IMPORT_VALIDATION_FAILED",
                "failed_rows": failed_count,
            }
            if failed_count
            else {}
        )
        locked.version += 1
        locked.updated_at = (
            datetime.now(
                timezone.utc
            ).replace(
                tzinfo=None
            )
        )
        await db.commit()
        return failed_count == 0
    except Exception:
        await db.rollback()
        raise
    finally:
        await _close(
            token,
            db,
        )


async def _import_rows(
    company_id: int,
    job_id: UUID,
) -> None:
    while True:
        token, db = await _tenant_session(
            company_id
        )
        try:
            job = await db.scalar(
                select(ProductImportJob)
                .where(
                    ProductImportJob.company_id
                    == int(company_id),
                    ProductImportJob.id
                    == job_id,
                )
                .with_for_update()
            )
            if job is None:
                raise ValueError(
                    "Product import job not found."
                )

            actor = await db.scalar(
                select(Driver).where(
                    Driver.company_id
                    == int(company_id),
                    Driver.id
                    == int(job.created_by),
                    Driver.is_active.is_(
                        True
                    ),
                )
            )
            if actor is None:
                raise ProductImportTerminalError(
                    "The import actor is no longer active."
                )

            access = InventoryAccess(
                db,
                actor,
            )
            try:
                for permission in (
                    "catalog.manage",
                    "catalog.publish",
                    "pricing.manage",
                ):
                    await access.require(
                        permission,
                        any_location=True,
                    )
            except HTTPException as exc:
                raise ProductImportTerminalError(
                    "The import actor no longer has the required permissions."
                ) from exc

            rows = list(
                (
                    await db.scalars(
                        select(ProductImportRow)
                        .where(
                            ProductImportRow.company_id
                            == int(company_id),
                            ProductImportRow.job_id
                            == job_id,
                            ProductImportRow.status
                            == "VALID",
                        )
                        .order_by(
                            ProductImportRow.row_number.asc()
                        )
                        .limit(
                            IMPORT_BATCH
                        )
                        .with_for_update(
                            skip_locked=True
                        )
                    )
                ).all()
            )

            if not rows:
                job.status = "COMPLETED"
                job.processed_rows = int(
                    job.valid_rows
                )
                job.finished_at = (
                    datetime.now(
                        timezone.utc
                    ).replace(
                        tzinfo=None
                    )
                )
                job.error_summary = {}
                job.version += 1

                await db.execute(
                    delete(
                        ProductImportRow
                    ).where(
                        ProductImportRow.company_id
                        == int(company_id),
                        ProductImportRow.job_id
                        == job_id,
                        ProductImportRow.status
                        == "IMPORTED",
                    )
                )
                await db.commit()
                return

            specs = []
            for row in rows:
                data = dict(
                    row.normalized_data
                    or {}
                )
                specs.append(
                    SimpleProductSpec(
                        name=str(
                            data["name"]
                        ),
                        family_name=data.get(
                            "family_name"
                        ),
                        package_uom_code=data.get(
                            "package_uom_code"
                        ),
                        units_per_package=int(
                            data[
                                "units_per_package"
                            ]
                        ),
                        package_price=(
                            Decimal(
                                str(
                                    data[
                                        "package_price"
                                    ]
                                )
                            )
                            if data.get(
                                "package_price"
                            )
                            is not None
                            else None
                        ),
                        unit_price=Decimal(
                            str(
                                data[
                                    "unit_price"
                                ]
                            )
                        ),
                        unit_barcode=data.get(
                            "unit_barcode"
                        ),
                        package_barcode=data.get(
                            "package_barcode"
                        ),
                        lot_control_mode=str(
                            data["lot_control_mode"]
                        ),
                        expiry_control_mode=str(
                            data["expiry_control_mode"]
                        ),
                    )
                )

            request_id = uuid5(
                job_id,
                (
                    "rows:"
                    f"{int(rows[0].row_number)}:"
                    f"{int(rows[-1].row_number)}"
                ),
            )
            created = await create_products_and_prices(
                db,
                actor=actor,
                request_id=request_id,
                specs=specs,
            )

            for (
                row,
                (variant, _prices),
            ) in zip(
                rows,
                created,
                strict=True,
            ):
                row.status = "IMPORTED"
                row.product_variant_id = int(
                    variant.id
                )
                row.version += 1

            imported = await db.scalar(
                select(
                    func.count(
                        ProductImportRow.id
                    )
                ).where(
                    ProductImportRow.company_id
                    == int(company_id),
                    ProductImportRow.job_id
                    == job_id,
                    ProductImportRow.status
                    == "IMPORTED",
                )
            )
            job.processed_rows = int(
                imported
                or 0
            )
            job.status = "IMPORTING"
            job.version += 1
            job.updated_at = (
                datetime.now(
                    timezone.utc
                ).replace(
                    tzinfo=None
                )
            )
            await db.commit()
        except Exception:
            await db.rollback()
            raise
        finally:
            await _close(
                token,
                db,
            )


async def run_product_import_job(
    *,
    company_id: int,
    job_id: UUID,
) -> None:
    (
        payload,
        file_name,
        status,
        _headers,
        _mapping,
    ) = await _load_job(
        company_id,
        job_id,
    )

    if status in {
        "COMPLETED",
        "VALIDATION_FAILED",
        "FAILED",
        "NEEDS_MAPPING",
    }:
        return

    if payload is not None:
        await _set_status(
            company_id,
            job_id,
            "PARSING",
            started_at=(
                datetime.now(
                    timezone.utc
                ).replace(
                    tzinfo=None
                )
            ),
            error_summary={},
        )
        headers, rows = await asyncio.to_thread(
            parse_source,
            file_name,
            payload,
        )
        suggestions = suggest_mapping(
            headers
        )
        await _stage_source(
            company_id,
            job_id,
            headers,
            rows,
            suggestions,
        )

    token, db = await _tenant_session(
        company_id
    )
    try:
        job = await db.scalar(
            select(ProductImportJob).where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
        )
        if job is None:
            raise ProductImportTerminalError(
                "Product import job not found."
            )
        status = str(job.status)
        has_source = (
            job.source_payload
            is not None
        )
    finally:
        await _close(
            token,
            db,
        )

    if status == "NEEDS_MAPPING":
        return

    if (
        status
        in {"QUEUED", "PARSING"}
        and not has_source
    ):
        raise ProductImportTerminalError(
            "Import state is inconsistent and cannot be resumed safely."
        )

    if status == "VALIDATING":
        if not await _validate_rows(
            company_id,
            job_id,
        ):
            return
        status = "IMPORTING"

    if status == "IMPORTING":
        await _import_rows(
            company_id,
            job_id,
        )
        return

    raise ProductImportTerminalError(
        f"Unsupported import state: {status}"
    )


async def mark_import_runtime_failure(
    *,
    company_id: int,
    job_id: UUID,
    message: str,
    final_attempt: bool,
    retryable: bool,
) -> None:
    token, db = await _tenant_session(
        company_id
    )
    try:
        job = await db.scalar(
            select(ProductImportJob)
            .where(
                ProductImportJob.company_id
                == int(company_id),
                ProductImportJob.id
                == job_id,
            )
            .with_for_update()
        )
        if job is None:
            return

        current_status = str(
            job.status
        )
        summary = {
            "code": (
                "PRODUCT_IMPORT_FAILED"
                if final_attempt
                else "PRODUCT_IMPORT_RETRYING"
            ),
            "technical": str(
                message
            )[:500],
            "retryable": (
                bool(retryable)
                if final_attempt
                else False
            ),
            "resume_status": current_status,
        }

        job.error_summary = summary
        if final_attempt:
            job.status = "FAILED"
            job.finished_at = (
                datetime.now(
                    timezone.utc
                ).replace(
                    tzinfo=None
                )
            )
        job.version += 1
        job.updated_at = (
            datetime.now(
                timezone.utc
            ).replace(
                tzinfo=None
            )
        )
        await db.commit()
    except Exception:
        await db.rollback()
    finally:
        await _close(
            token,
            db,
        )
