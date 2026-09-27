from __future__ import annotations

import asyncio
import sys
from io import BytesIO
from pathlib import Path

from sqlalchemy import text


ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "wa_backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

checks = 0
failures: list[str] = []


def check(condition: bool, label: str) -> None:
    global checks
    checks += 1
    print(("[PASS] " if condition else "[FAIL] ") + label)
    if not condition:
        failures.append(label)


def static_checks() -> None:
    from openpyxl import load_workbook

    from domains.simple_products.imports.infrastructure.template import (
        build_product_import_template,
    )
    from domains.simple_products.imports.domain.mapping import (
        suggest_mapping,
    )
    from domains.simple_products.imports.domain.normalization import (
        normalize_raw_row,
    )

    inherited = normalize_raw_row(
        {
            "Product": "General Item",
            "Price": "1.250",
        },
        {
            "name": "Product",
            "unit_price": "Price",
        },
        default_lot_control_mode="NONE",
        default_expiry_control_mode="OPTIONAL",
    )
    check(
        inherited["lot_control_mode"] == "NONE"
        and inherited["expiry_control_mode"] == "OPTIONAL",
        "Rows without tracking columns inherit the import snapshot",
    )

    overridden = normalize_raw_row(
        {
            "Product": "Food Item",
            "Price": "1.250",
            "Lot Tracking": "REQUIRED",
            "Expiry Tracking": "REQUIRED",
        },
        {
            "name": "Product",
            "unit_price": "Price",
            "lot_control_mode": "Lot Tracking",
            "expiry_control_mode": "Expiry Tracking",
        },
        default_lot_control_mode="NONE",
        default_expiry_control_mode="NONE",
    )
    check(
        overridden["lot_control_mode"] == "REQUIRED"
        and overridden["expiry_control_mode"] == "REQUIRED",
        "Non-empty row tracking values override import defaults",
    )

    localized = normalize_raw_row(
        {
            "Product": "Mixed Item",
            "Price": "1.250",
            "Lot Tracking": "لا",
            "Expiry Tracking": "إلزامي",
        },
        {
            "name": "Product",
            "unit_price": "Price",
            "lot_control_mode": "Lot Tracking",
            "expiry_control_mode": "Expiry Tracking",
        },
        default_lot_control_mode="REQUIRED",
        default_expiry_control_mode="NONE",
    )
    check(
        localized["lot_control_mode"] == "NONE"
        and localized["expiry_control_mode"] == "REQUIRED",
        "Localized row tracking values normalize to canonical API codes",
    )

    explicit_default = normalize_raw_row(
        {
            "Product": "Defaulted Item",
            "Price": "1.250",
            "Lot Tracking": "استخدام الافتراضي",
            "Expiry Tracking": "استخدام الافتراضي",
        },
        {
            "name": "Product",
            "unit_price": "Price",
            "lot_control_mode": "Lot Tracking",
            "expiry_control_mode": "Expiry Tracking",
        },
        default_lot_control_mode="OPTIONAL",
        default_expiry_control_mode="NONE",
    )
    check(
        explicit_default["lot_control_mode"] == "OPTIONAL"
        and explicit_default["expiry_control_mode"] == "NONE",
        "Explicit Use-default tracking cells resolve to the import snapshot",
    )

    template_bytes = build_product_import_template(
        locale="ar",
    )
    workbook = load_workbook(
        BytesIO(template_bytes),
    )
    template_sheet = workbook["المنتجات"]
    validations = list(
        template_sheet.data_validations.dataValidation
    )
    validation_ranges = {
        str(cell_range)
        for validation in validations
        for cell_range in validation.ranges.ranges
    }
    list_sheet = workbook["_wanasah_lists"]
    check(
        len(validations) == 2
        and "I2:I50001" in validation_ranges
        and "J2:J50001" in validation_ranges
        and list_sheet.sheet_state == "hidden"
        and [
            list_sheet.cell(row=row, column=1).value
            for row in range(1, 5)
        ]
        == [
            "استخدام الافتراضي",
            "بدون",
            "اختياري",
            "إلزامي",
        ],
        "Excel template provides validated tracking dropdowns for 50,000 product rows",
    )
    check(
        template_sheet["I1"].comment is not None
        and template_sheet["J1"].comment is not None
        and template_sheet.freeze_panes == "A2",
        "Excel template explains blank/default behavior at the tracking columns",
    )

    suggestions = suggest_mapping(
        [
            "Product",
            "Unit Price",
            "Lot Tracking",
            "Expiry Tracking",
        ]
    )
    check(
        suggestions.get("lot_control_mode") == "Lot Tracking"
        and suggestions.get("expiry_control_mode")
        == "Expiry Tracking",
        "Import header detection recognizes tracking columns",
    )

    api = (
        BACKEND / "domains/simple_products/imports/api/router.py"
    ).read_text(encoding="utf-8")
    worker = (
        BACKEND / "domains/simple_products/imports/application/worker.py"
    ).read_text(encoding="utf-8")
    normalization = (
        BACKEND / "domains/simple_products/imports/domain/normalization.py"
    ).read_text(encoding="utf-8")
    execution = (
        BACKEND / "domains/simple_products/imports/application/execution_service.py"
    ).read_text(encoding="utf-8")
    localization = (
        BACKEND / "domains/simple_products/imports/domain/localization.py"
    ).read_text(encoding="utf-8")
    queue = (
        BACKEND / "domains/simple_products/imports/infrastructure/queue.py"
    ).read_text(encoding="utf-8")
    models = (
        BACKEND / "models.py"
    ).read_text(encoding="utf-8")
    migration = (
        BACKEND
        / "alembic"
        / "versions"
        / "f3c8a1d4e6b2_product_import_tracking_defaults.py"
    ).read_text(encoding="utf-8")
    repair_migration = (
        BACKEND
        / "alembic"
        / "versions"
        / "e7a1c4d9b2f6_import_tracking_constraints_repair.py"
    ).read_text(encoding="utf-8")
    import_upload = (
        ROOT
        / "dashboard/src/pages/products/import/useImportProductUpload.ts"
    ).read_text(encoding="utf-8")
    import_modal = (
        ROOT
        / "dashboard/src/pages/products/import/ImportProductModal.tsx"
    ).read_text(encoding="utf-8")
    import_start = (
        ROOT
        / "dashboard/src/pages/products/import/ImportProductStartPanel.tsx"
    ).read_text(encoding="utf-8")
    import_mapping = (
        ROOT
        / "dashboard/src/pages/products/import/ImportProductMappingPanel.tsx"
    ).read_text(encoding="utf-8")
    import_downloads = (
        ROOT
        / "dashboard/src/pages/products/import/createImportDownloads.ts"
    ).read_text(encoding="utf-8")
    import_quick_guide = (
        ROOT
        / "dashboard/src/pages/products/import/ImportProductQuickGuide.tsx"
    ).read_text(encoding="utf-8")
    import_template = (
        BACKEND / "domains/simple_products/imports/infrastructure/template.py"
    ).read_text(encoding="utf-8")
    controls = (
        ROOT
        / "dashboard/src/pages/products/tracking/ProductTrackingFields.tsx"
    ).read_text(encoding="utf-8")
    mode_picker = (
        ROOT
        / "dashboard/src/pages/products/tracking/ProductTrackingModePicker.tsx"
    ).read_text(encoding="utf-8")

    check(
        "default_lot_control_mode: str | None = Form(None)" in api
        and "default_expiry_control_mode: str | None = Form(None)" in api
        and "resolve_product_tracking_modes(" in api,
        "Import API resolves explicit or company tracking defaults",
    )
    check(
        '"default_lot_control_mode"' in api
        and '"default_expiry_control_mode"' in api,
        "Import responses expose the immutable tracking snapshot",
    )
    check(
        '"lot_control_mode"' in execution
        and '"expiry_control_mode"' in execution
        and "lot_control_mode=str(" in execution
        and "expiry_control_mode=str(" in execution
        and "default_lot_control_mode" in normalization
        and "default_expiry_control_mode" in normalization,
        "Normalization and execution preserve Product tracking contract",
    )
    check(
        "canonical_tracking_value" in normalization
        and "tracking_value_aliases" in localization
        and "IMPORT_TRACKING_DEFAULT_SENTINEL" in localization
        and '"استخدام الافتراضي"' in localization
        and '"use default"' in localization
        and '"لا": "NONE"' in localization
        and '"اختياري": "OPTIONAL"' in localization
        and '"إلزامي": "REQUIRED"' in localization
        and '"required": "REQUIRED"' in localization,
        "Import accepts localized tracking values and explicit Use-default without changing stored codes",
    )
    check(
        "default_lot_control_mode" in queue
        and "default_expiry_control_mode" in queue
        and "existing_lot_default" in queue
        and "existing_expiry_default" in queue,
        "Import replay identity includes both tracking defaults",
    )
    check(
        "default_lot_control_mode = Column(String(20), nullable=False)"
        in models
        and "default_expiry_control_mode = Column(String(20), nullable=False)"
        in models,
        "Import jobs persist an explicit tracking snapshot",
    )
    check(
        "import_job_lot_mode" in models
        and "import_job_expiry_mode" in models,
        "ORM constrains import tracking snapshot values",
    )
    check(
        "down_revision = \"d4a7c9e2f1b5\"" in migration
        and "SET\n            default_lot_control_mode = 'REQUIRED'" in migration
        and "default_expiry_control_mode = 'REQUIRED'" in migration,
        "Migration preserves historical REQUIRED/REQUIRED import behavior",
    )
    check(
        '"default_lot_control_mode"' in import_upload
        and '"default_expiry_control_mode"' in import_upload
        and "lotControlMode" in import_start
        and "expiryControlMode" in import_start
        and "ProductTrackingFields" in import_start
        and "IMPORT_MAPPING_FIELDS.map" in import_mapping
        and "onMappingChange" in import_mapping,
        "Products UI sends import defaults and supports row mapping",
    )
    check(
        '"/import-template"' in api
        and "build_product_import_template" in api
        and "DataValidation(" in import_template
        and "MAX_TEMPLATE_ROWS = 50_000" in import_template
        and '"products.importGuideUseDefault"' in import_quick_guide
        and '"products.importGuideBulkTip"' in import_quick_guide
        and '"products.importTrackingValueHint"' in import_start
        and "/simple-products/import-template?locale=" in import_downloads,
        "Import template and visible guidance expose simple bulk tracking defaults",
    )
    check(
        "ProductTrackingModePicker" in controls
        and '"NONE"' in mode_picker
        and '"OPTIONAL"' in mode_picker
        and '"REQUIRED"' in mode_picker
        and "products.tracking.lotModes" in mode_picker
        and "products.tracking.expiryModes" in mode_picker
        and "products.tracking.lotExample" in mode_picker
        and "products.tracking.expiryExample" in mode_picker,
        "Tracking controls keep API codes language-neutral while using contextual labels",
    )
    check(
        'down_revision = "f3c8a1d4e6b2"' in repair_migration
        and "ck_product_import_jobs_import_job_lot_mode"
        in repair_migration
        and "ck_product_import_jobs_import_job_expiry_mode"
        in repair_migration
        and "_column_constraints" in repair_migration
        and "_validate_existing_constraint" in repair_migration
        and "_canonicalize_constraint" in repair_migration,
        "Follow-up migration detects, validates and canonicalizes import tracking constraints",
    )


async def database_checks() -> None:
    from database import engine

    async with engine.connect() as conn:
        columns = {
            str(row.column_name): str(row.is_nullable)
            for row in (
                await conn.execute(
                    text(
                        """
                        SELECT column_name, is_nullable
                        FROM information_schema.columns
                        WHERE table_schema = current_schema()
                          AND table_name = 'product_import_jobs'
                          AND column_name IN (
                              'default_lot_control_mode',
                              'default_expiry_control_mode'
                          )
                        """
                    )
                )
            ).all()
        }
        check(
            columns
            == {
                "default_lot_control_mode": "NO",
                "default_expiry_control_mode": "NO",
            },
            "Database stores both tracking snapshots as NOT NULL",
        )

        constraints = set(
            (
                await conn.execute(
                    text(
                        """
                        SELECT conname
                        FROM pg_constraint
                        WHERE conrelid = 'product_import_jobs'::regclass
                          AND contype = 'c'
                          AND conname IN (
                            'ck_product_import_jobs_import_job_lot_mode',
                            'ck_product_import_jobs_import_job_expiry_mode'
                          )
                        """
                    )
                )
            ).scalars().all()
        )
        check(
            constraints
            == {
                "ck_product_import_jobs_import_job_lot_mode",
                "ck_product_import_jobs_import_job_expiry_mode",
            },
            "Database constrains import tracking snapshot values",
        )

        rls = (
            await conn.execute(
                text(
                    """
                    SELECT relrowsecurity, relforcerowsecurity
                    FROM pg_class
                    WHERE relname = 'product_import_jobs'
                    """
                )
            )
        ).one_or_none()
        check(
            rls is not None
            and bool(rls.relrowsecurity)
            and bool(rls.relforcerowsecurity),
            "Product import jobs keep ENABLE + FORCE RLS",
        )

        policies = (
            await conn.execute(
                text(
                    """
                    SELECT qual, with_check
                    FROM pg_policies
                    WHERE schemaname = current_schema()
                      AND tablename = 'product_import_jobs'
                    """
                )
            )
        ).all()
        check(
            any(
                "app.current_tenant" in str(row.qual or "")
                and "company_id" in str(row.qual or "")
                and "app.current_tenant"
                in str(row.with_check or "")
                and "company_id"
                in str(row.with_check or "")
                for row in policies
            ),
            "Product import jobs remain tenant isolated by RLS",
        )

    await engine.dispose()


def main() -> int:
    try:
        static_checks()
        asyncio.run(database_checks())
    except Exception as exc:
        check(False, "Gate completed without unexpected exception")
        print(f"UNEXPECTED_EXCEPTION={exc!r}")

    print()
    print(f"CHECKS={checks}")
    print(f"FAILURES={len(failures)}")
    for failure in failures:
        print(f"FAIL: {failure}")

    if failures:
        print("PRODUCT_IMPORT_TRACKING_GATE=FAIL")
        return 1

    print("PRODUCT_IMPORT_TRACKING_GATE=PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
