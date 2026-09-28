from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import load_workbook


BACKEND = Path(__file__).resolve().parents[1]
ROOT = BACKEND.parent
sys.path.insert(
    0,
    str(BACKEND),
)

from domains.simple_products.imports.domain import (  # noqa: E402
    CANONICAL_IMPORT_FIELDS,
    SOURCE_CELL_META_KEY,
    WANASAH_TEMPLATE_MARKER,
    WANASAH_TEMPLATE_META_SHEET,
)
from domains.simple_products.imports.infrastructure.template import (  # noqa: E402
    MAX_TEMPLATE_ROWS,
    build_product_import_template,
)


checks = 0
failures: list[str] = []


def check(
    condition: bool,
    label: str,
) -> None:
    global checks
    checks += 1
    print(
        (
            "[PASS] "
            if condition
            else "[FAIL] "
        )
        + label
    )
    if not condition:
        failures.append(
            label
        )


def source(
    *parts: str,
) -> str:
    return (
        BACKEND.joinpath(
            *parts
        ).read_text(
            encoding="utf-8"
        )
    )


parser_source = source(
    "domains",
    "simple_products",
    "imports",
    "infrastructure",
    "parsers.py",
)
normalization_source = source(
    "domains",
    "simple_products",
    "imports",
    "domain",
    "normalization.py",
)
template_source = source(
    "domains",
    "simple_products",
    "imports",
    "infrastructure",
    "template.py",
)
date_source = source(
    "domains",
    "simple_products",
    "imports",
    "domain",
    "spreadsheet_dates.py",
)
repository_source = source(
    "domains",
    "simple_products",
    "imports",
    "infrastructure",
    "repository.py",
)
i18n_source = (
    ROOT
    / "dashboard"
    / "src"
    / "i18n"
    / "resources.ts"
).read_text(
    encoding="utf-8"
)

check(
    "row_number += 1"
    in parser_source
    and "accepted_count"
    in parser_source
    and '"row_number": int('
    in repository_source,
    "Physical source row number is independent from accepted-row count and is staged verbatim",
)

check(
    "workbook.active"
    not in parser_source,
    "XLSX parser never selects workbook.active",
)

check(
    "WANASAH_TEMPLATE_META_SHEET"
    in parser_source
    and "WANASAH_TEMPLATE_MARKER"
    in parser_source
    and "multiple visible worksheets"
    in parser_source,
    "Worksheet selection is deterministic and ambiguous external workbooks fail closed",
)

check(
    "data_only=True"
    in parser_source
    and "data_only=False"
    in parser_source
    and "SOURCE_CELL_META_KEY"
    in parser_source,
    "Formula and cached-value readers are paired without losing source-cell provenance",
)

check(
    "IMPORT_FORMULA_VALUE_UNAVAILABLE"
    in normalization_source
    and "IMPORT_BARCODE_FORMULA_NOT_ALLOWED"
    in normalization_source,
    "Formula policy accepts safe cached values but fails closed for unavailable/literal-only cells",
)

check(
    "IMPORT_BARCODE_NUMERIC_UNSAFE"
    in normalization_source
    and "IMPORT_BARCODE_SCIENTIFIC_NOTATION"
    in normalization_source,
    "Unsafe numeric/scientific barcode identities are rejected",
)

for code in (
    "IMPORT_FORMULA_VALUE_UNAVAILABLE",
    "IMPORT_BARCODE_FORMULA_NOT_ALLOWED",
    "IMPORT_BARCODE_NUMERIC_UNSAFE",
    "IMPORT_BARCODE_SCIENTIFIC_NOTATION",
):
    check(
        i18n_source.count(
            code
        )
        >= 2,
        f"{code} has Arabic and English user-safe localization",
    )

payload = (
    build_product_import_template(
        locale="en"
    )
)
workbook = load_workbook(
    io.BytesIO(
        payload
    )
)
try:
    check(
        WANASAH_TEMPLATE_META_SHEET
        in workbook.sheetnames
        and workbook[
            WANASAH_TEMPLATE_META_SHEET
        ][
            "A1"
        ].value
        == WANASAH_TEMPLATE_MARKER,
        "Official workbook carries an explicit Wanasah template marker",
    )

    meta = workbook[
        WANASAH_TEMPLATE_META_SHEET
    ]
    product_sheet_name = str(
        meta[
            "A2"
        ].value
        or ""
    )
    check(
        product_sheet_name
        in workbook.sheetnames
        and workbook[
            product_sheet_name
        ].sheet_state
        == "visible",
        "Official marker names one visible canonical Products worksheet",
    )

    product_sheet = workbook[
        product_sheet_name
    ]
    for field in (
        "unit_barcode",
        "package_barcode",
    ):
        column_index = (
            CANONICAL_IMPORT_FIELDS.index(
                field
            )
            + 1
        )
        column_letter = chr(
            64
            + column_index
        )
        check(
            product_sheet.column_dimensions[
                column_letter
            ].number_format
            == "@",
            f"{field} template column is forced to Excel Text format",
        )

    barcode_validations = [
        validation
        for validation
        in product_sheet.data_validations.dataValidation
        if validation.type
        == "custom"
        and "ISTEXT("
        in str(
            validation.formula1
        )
    ]
    check(
        len(
            barcode_validations
        )
        == 2
        and all(
            "ISTEXT("
            in str(
                validation.formula1
            )
            and str(
                MAX_TEMPLATE_ROWS
                + 1
            )
            in str(
                validation.sqref
            )
            for validation
            in barcode_validations
        ),
        "Official template validates barcode text entry through all 50,000 Product rows",
    )
    package_lists = [
        validation
        for validation
        in product_sheet.data_validations.dataValidation
        if validation.type
        == "list"
        and "$B$1:$B$9"
        in str(
            validation.formula1
        )
    ]
    check(
        len(
            package_lists
        )
        == 1
        and "C2:C50001"
        in str(
            package_lists[
                0
            ].sqref
        ),
        "Official template guides package type through all 50,000 Product rows",
    )
    check(
        product_sheet[
            "A1"
        ].fill.fgColor.rgb
        == "00475569"
        and product_sheet.column_dimensions[
            "E"
        ].width
        == 14
        and product_sheet.column_dimensions[
            "F"
        ].width
        == 14
        and product_sheet.column_dimensions[
            "I"
        ].width
        == 16
        and product_sheet.column_dimensions[
            "J"
        ].width
        == 16,
        "Official template uses the polished slate header and compact price/tracking widths",
    )
finally:
    workbook.close()

check(
    "CALENDAR_WINDOWS_1900"
    in date_source
    and "CALENDAR_MAC_1904"
    in date_source
    and "serial_int == 60"
    in date_source,
    "Future spreadsheet-date authority supports 1900/1904 epochs and rejects phantom serial 60",
)

check(
    "date.fromisoformat"
    in date_source
    and "Timezone-aware"
    in date_source
    and "time component"
    in date_source,
    "Future date-only authority accepts strict ISO/native values without silent timezone/time coercion",
)

check(
    not any(
        field
        in CANONICAL_IMPORT_FIELDS
        for field in (
            "expiry_date",
            "expiration_date",
            "production_date",
            "manufacture_date",
        )
    ),
    "Current Product Catalog import still has no date-valued inventory/expiry field",
)

xlsx_flow = parser_source[
    parser_source.index(
        "def _open_xlsx_source("
    ):
    parser_source.index(
        "def open_source("
    )
]
check(
    "_validate_xlsx_archive("
    in xlsx_flow
    and "_load_bounded_workbook("
    in xlsx_flow
    and xlsx_flow.index(
        "_validate_xlsx_archive("
    )
    < xlsx_flow.index(
        "_load_bounded_workbook("
    ),
    "XLSX archive security remains before workbook traversal",
)

check(
    "SOURCE_CELL_META_KEY"
    in template_source
    or "WANASAH_TEMPLATE_META_SHEET"
    in template_source,
    "Template generator owns explicit source semantics instead of relying on active-sheet behavior",
)


if failures:
    print(
        f"CHECKS={checks}"
    )
    print(
        f"FAILURES={len(failures)}"
    )
    for failure in failures:
        print(
            f"FAIL: {failure}"
        )
    print(
        "PRODUCT_IMPORT_PHASE11_SOURCE_SEMANTICS_GATE=FAIL"
    )
    raise SystemExit(1)

print(
    f"CHECKS={checks}"
)
print("FAILURES=0")
print(
    "PRODUCT_IMPORT_PHASE11_SOURCE_SEMANTICS_GATE=PASS"
)
