from __future__ import annotations

import asyncio
import gc
import io
import sys
import tracemalloc
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "wa_backend"
sys.path.insert(
    0,
    str(BACKEND),
)

from domains.simple_products.imports.infrastructure.parsers import (  # noqa: E402
    MAX_IMPORT_ROWS,
    _DiskBackedSharedStrings,
    open_source,
)
from domains.simple_products.imports.infrastructure.repository import (  # noqa: E402
    insert_staged_rows,
)


BATCH_SIZE = 1_000
MEMORY_SLACK_BYTES = (
    768 * 1024
)

checks = 0
failures: list[str] = []


def check(
    condition: bool,
    label: str,
    detail: str = "",
) -> None:
    global checks
    checks += 1
    prefix = (
        "[PASS]"
        if condition
        else "[FAIL]"
    )
    suffix = (
        f" — {detail}"
        if detail
        else ""
    )
    print(
        f"{prefix} {label}{suffix}"
    )
    if not condition:
        failures.append(label)


class FakeDb:
    def __init__(
        self,
    ) -> None:
        self.max_batch = 0
        self.calls = 0

    async def execute(
        self,
        _statement,
        parameters,
    ) -> None:
        size = len(parameters)
        self.max_batch = max(
            self.max_batch,
            size,
        )
        self.calls += 1


def csv_payload(
    count: int,
) -> bytes:
    buffer = io.BytesIO()
    buffer.write(
        b"Product,Unit Price\n"
    )
    for index in range(
        count
    ):
        buffer.write(
            (
                f"P{index},1.000\n"
            ).encode("utf-8")
        )
    return buffer.getvalue()


def xlsx_payload(
    count: int,
) -> bytes:
    buffer = io.BytesIO()
    workbook = Workbook(
        write_only=True
    )
    sheet = workbook.create_sheet(
        "Products"
    )
    sheet.append(
        [
            "Product",
            "Unit Price",
        ]
    )
    for index in range(
        count
    ):
        sheet.append(
            [
                f"Product-{index:05d}",
                "1.000",
            ]
        )
    workbook.save(
        buffer
    )
    workbook.close()
    return buffer.getvalue()


def shared_strings_xml(
    count: int,
) -> bytes:
    buffer = io.BytesIO()
    buffer.write(
        (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<sst xmlns="http://schemas.openxmlformats.org/'
            'spreadsheetml/2006/main">'
        ).encode("utf-8")
    )
    for index in range(
        count
    ):
        buffer.write(
            (
                f"<si><t>S{index}</t></si>"
            ).encode("utf-8")
        )
    buffer.write(
        b"</sst>"
    )
    return buffer.getvalue()


async def measure_xlsx_peak(
    count: int,
) -> tuple[int, int, int]:
    payload = xlsx_payload(
        count
    )
    db = FakeDb()

    # Source bytes/workbook bootstrap are outside the incremental working-set
    # measurement, matching the CSV gate and Phase 12 SourceStore boundary.
    with open_source(
        "products.xlsx",
        payload,
    ) as source:
        gc.collect()
        tracemalloc.start()
        try:
            total = await insert_staged_rows(
                db,
                company_id=1,
                job_id=uuid4(),
                rows=source.rows,
                batch_size=BATCH_SIZE,
                max_rows=MAX_IMPORT_ROWS,
            )
            _current, peak = (
                tracemalloc.get_traced_memory()
            )
        finally:
            tracemalloc.stop()

    return (
        total,
        int(
            peak
        ),
        int(
            db.max_batch
        ),
    )


def measure_shared_strings_peak(
    count: int,
) -> tuple[int, int]:
    payload = shared_strings_xml(
        count
    )
    store = (
        _DiskBackedSharedStrings()
    )
    gc.collect()
    tracemalloc.start()
    try:
        store.load(
            io.BytesIO(
                payload
            )
        )
        self_check = (
            store[0]
            == "S0"
            and store[
                count - 1
            ]
            == f"S{count - 1}"
        )
        _current, peak = (
            tracemalloc.get_traced_memory()
        )
    finally:
        tracemalloc.stop()
        store.close()

    if not self_check:
        raise RuntimeError(
            "Disk-backed shared-string lookup failed."
        )
    return (
        count,
        int(peak),
    )


async def measure_peak(
    count: int,
) -> tuple[int, int, int]:
    payload = csv_payload(
        count
    )
    db = FakeDb()

    # Source bytes and the already-open immutable source wrapper are excluded
    # deliberately: Phase 4 measures the incremental parser/staging working
    # set. SourceStore/blob memory is owned by Phase 12.
    with open_source(
        "products.csv",
        payload,
    ) as source:
        gc.collect()
        tracemalloc.start()
        try:
            total = await insert_staged_rows(
                db,
                company_id=1,
                job_id=uuid4(),
                rows=source.rows,
                batch_size=BATCH_SIZE,
                max_rows=MAX_IMPORT_ROWS,
            )
            _current, peak = (
                tracemalloc.get_traced_memory()
            )
        finally:
            tracemalloc.stop()

    return (
        total,
        int(peak),
        int(db.max_batch),
    )


async def main() -> None:
    parser_source = (
        BACKEND
        / "domains"
        / "simple_products"
        / "imports"
        / "infrastructure"
        / "parsers.py"
    ).read_text(
        encoding="utf-8"
    )
    repository_source = (
        BACKEND
        / "domains"
        / "simple_products"
        / "imports"
        / "infrastructure"
        / "repository.py"
    ).read_text(
        encoding="utf-8"
    )

    check(
        "rows.append(" not in parser_source
        and "list[dict[str, Any]]"
        not in parser_source,
        "Parser does not materialize an import-sized Product row list",
    )
    check(
        "yield parsed" in parser_source
        and "ParsedRow" in parser_source,
        "Parser contract yields Product rows incrementally with source row numbers",
    )
    xlsx_open_source = parser_source[
        parser_source.index(
            "def _open_xlsx_source("
        ):
    ]
    check(
        "read_only=True"
        in parser_source
        and xlsx_open_source.index(
            "_validate_xlsx_archive("
        )
        < xlsx_open_source.index(
            "_load_bounded_workbook("
        ),
        "XLSX archive security remains ahead of workbook traversal",
    )
    check(
        "_DiskBackedSharedStrings"
        in parser_source
        and "sqlite3.connect("
        in parser_source
        and "load_workbook("
        not in parser_source,
        "XLSX shared strings use bounded disk-backed storage instead of OpenPyXL list materialization",
    )
    check(
        "batch.clear()"
        in repository_source
        and "max_rows"
        in repository_source,
        "Repository flushes fixed staging batches and enforces a defensive row ceiling",
    )

    # Warm SQLAlchemy mapper/statement construction before measuring.
    await measure_peak(100)

    measurements: dict[
        int,
        tuple[int, int, int],
    ] = {}
    for count in (
        1_000,
        10_000,
        50_000,
    ):
        measurements[count] = (
            await measure_peak(
                count
            )
        )

    for (
        count,
        (
            total,
            peak,
            max_batch,
        ),
    ) in measurements.items():
        check(
            total == count,
            f"{count:,}-row stream stages every row",
            f"total={total}",
        )
        check(
            max_batch
            <= BATCH_SIZE,
            f"{count:,}-row stream never exceeds configured staging batch",
            f"max_batch={max_batch}",
        )
        print(
            "MEMORY_PEAK_"
            f"{count}="
            f"{peak}"
        )

    baseline_peak = (
        measurements[
            1_000
        ][1]
    )
    ten_k_peak = (
        measurements[
            10_000
        ][1]
    )
    fifty_k_peak = (
        measurements[
            50_000
        ][1]
    )
    allowed_peak = (
        baseline_peak
        + MEMORY_SLACK_BYTES
    )

    check(
        ten_k_peak
        <= allowed_peak,
        "10,000-row incremental working set stays within fixed memory envelope",
        (
            f"baseline={baseline_peak} "
            f"peak={ten_k_peak} "
            f"limit={allowed_peak}"
        ),
    )
    check(
        fifty_k_peak
        <= allowed_peak,
        "50,000-row incremental working set stays within the same fixed memory envelope as 1,000 rows",
        (
            f"baseline={baseline_peak} "
            f"peak={fifty_k_peak} "
            f"limit={allowed_peak}"
        ),
    )

    xlsx_1k = await measure_xlsx_peak(
        1_000
    )
    xlsx_50k = await measure_xlsx_peak(
        50_000
    )
    print(
        "XLSX_MEMORY_PEAK_1000="
        f"{xlsx_1k[1]}"
    )
    print(
        "XLSX_MEMORY_PEAK_50000="
        f"{xlsx_50k[1]}"
    )
    check(
        xlsx_1k[0] == 1_000
        and xlsx_50k[0]
        == 50_000,
        "50,000-row XLSX stream stages every row",
        (
            f"baseline_total={xlsx_1k[0]} "
            f"large_total={xlsx_50k[0]}"
        ),
    )
    check(
        xlsx_1k[2]
        <= BATCH_SIZE
        and xlsx_50k[2]
        <= BATCH_SIZE,
        "50,000-row XLSX stream never exceeds configured staging batch",
        (
            f"baseline_batch={xlsx_1k[2]} "
            f"large_batch={xlsx_50k[2]}"
        ),
    )
    xlsx_memory_budget = (
        16 * 1024 * 1024
    )
    check(
        xlsx_50k[1]
        <= xlsx_memory_budget,
        "50,000-row XLSX incremental working set stays inside explicit 16 MiB budget",
        (
            f"baseline={xlsx_1k[1]} "
            f"peak={xlsx_50k[1]} "
            f"limit={xlsx_memory_budget}"
        ),
    )

    shared_1k_count, shared_1k_peak = (
        measure_shared_strings_peak(
            1_000
        )
    )
    shared_50k_count, shared_50k_peak = (
        measure_shared_strings_peak(
            50_000
        )
    )
    shared_allowed_peak = (
        shared_1k_peak
        + (3 * 1024 * 1024)
    )
    print(
        "SHARED_STRINGS_MEMORY_PEAK_1000="
        f"{shared_1k_peak}"
    )
    print(
        "SHARED_STRINGS_MEMORY_PEAK_50000="
        f"{shared_50k_peak}"
    )
    check(
        shared_1k_count == 1_000
        and shared_50k_count
        == 50_000,
        "Disk-backed XLSX shared-string store indexes every entry",
    )
    check(
        shared_50k_peak
        <= shared_allowed_peak,
        "50,000 shared XLSX strings remain inside a fixed-memory envelope",
        (
            f"baseline={shared_1k_peak} "
            f"peak={shared_50k_peak} "
            f"limit={shared_allowed_peak}"
        ),
    )

    if failures:
        print(
            f"CHECKS={checks}"
        )
        print(
            f"FAILURES={len(failures)}"
        )
        for failure in failures:
            print(
                f"FAIL: {failure}"
            )
        print(
            "PRODUCT_IMPORT_PHASE4_STREAMING_GATE=FAIL"
        )
        raise SystemExit(1)

    print(
        f"CHECKS={checks}"
    )
    print("FAILURES=0")
    print(
        "PRODUCT_IMPORT_PHASE4_STREAMING_GATE=PASS"
    )


if __name__ == "__main__":
    asyncio.run(main())
