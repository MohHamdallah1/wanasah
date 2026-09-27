"""CSV/XLSX parsing infrastructure for Product Import.

Behavior is intentionally preserved from the legacy root worker during the
Phase 1 structural refactor.
"""
from __future__ import annotations

import csv
import io
import zipfile
from datetime import datetime
from typing import Any

from openpyxl import load_workbook

from domains.simple_products.imports.domain import (
    ProductImportTerminalError,
    normalize_import_token,
)


MAX_IMPORT_ROWS = 50_000
MAX_IMPORT_COLUMNS = 100
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_XLSX_ARCHIVE_ENTRIES = 500
MAX_XLSX_COMPRESSION_RATIO = 200


def _json_cell(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    clean = str(value).strip()
    return clean or None


def _header_layout(
    values: list[Any] | tuple[Any, ...],
) -> list[tuple[int, str]]:
    layout = [
        (
            index,
            str(value or "").strip(),
        )
        for index, value in enumerate(values)
        if str(value or "").strip()
    ]
    if not layout:
        raise ProductImportTerminalError(
            "The header row is empty."
        )
    if len(layout) > MAX_IMPORT_COLUMNS:
        raise ProductImportTerminalError(
            f"The file has more than {MAX_IMPORT_COLUMNS} columns."
        )

    normalized = [
        normalize_import_token(header)
        for _, header in layout
    ]
    if len(normalized) != len(set(normalized)):
        raise ProductImportTerminalError(
            "The header row contains duplicate columns after normalization."
        )
    return layout


def _row_from_layout(
    values: list[Any] | tuple[Any, ...],
    layout: list[tuple[int, str]],
) -> dict[str, Any]:
    return {
        header: _json_cell(
            values[index]
            if index < len(values)
            else None
        )
        for index, header in layout
    }


def _parse_csv(
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    decoded = None
    last_error = None
    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1256",
    ):
        try:
            decoded = payload.decode(encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc

    if decoded is None:
        raise ProductImportTerminalError(
            "CSV encoding is not supported."
        ) from last_error

    try:
        dialect = csv.Sniffer().sniff(
            decoded[:8192],
            delimiters=",;\t|",
        )
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(
        io.StringIO(decoded),
        dialect=dialect,
    )
    first = next(reader, None)
    if first is None:
        raise ProductImportTerminalError(
            "The file is empty."
        )

    layout = _header_layout(first)
    headers = [
        header
        for _, header in layout
    ]
    rows: list[dict[str, Any]] = []
    for values in reader:
        raw = _row_from_layout(
            values,
            layout,
        )
        if not any(
            value is not None
            for value in raw.values()
        ):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise ProductImportTerminalError(
                f"The import exceeds the {MAX_IMPORT_ROWS:,}-row safety limit."
            )
        rows.append(raw)

    if not rows:
        raise ProductImportTerminalError(
            "The file contains no product rows."
        )

    return headers, rows


def _validate_xlsx_archive(
    payload: bytes,
) -> None:
    try:
        with zipfile.ZipFile(
            io.BytesIO(payload)
        ) as archive:
            infos = archive.infolist()
            if not infos:
                raise ProductImportTerminalError(
                    "The Excel file is empty or invalid."
                )
            if (
                len(infos)
                > MAX_XLSX_ARCHIVE_ENTRIES
            ):
                raise ProductImportTerminalError(
                    "The Excel archive has an abnormal number of internal entries."
                )

            total = 0
            for info in infos:
                total += int(info.file_size)
                if (
                    total
                    > MAX_XLSX_UNCOMPRESSED_BYTES
                ):
                    raise ProductImportTerminalError(
                        "The uncompressed Excel content exceeds the safety limit."
                    )
                if (
                    info.compress_size > 0
                    and info.file_size > 1_000_000
                    and (
                        info.file_size
                        / info.compress_size
                    )
                    > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise ProductImportTerminalError(
                        "The Excel file was rejected because of an abnormal compression ratio."
                    )
                if info.filename.lower().endswith(
                    "vbaproject.bin"
                ):
                    raise ProductImportTerminalError(
                        "Macro-enabled Excel files are not supported."
                    )
    except zipfile.BadZipFile as exc:
        raise ProductImportTerminalError(
            "The Excel file is invalid or corrupted."
        ) from exc


def _parse_xlsx(
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    _validate_xlsx_archive(payload)
    try:
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise ProductImportTerminalError(
            "The Excel file could not be opened."
        ) from exc

    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(
            values_only=True
        )
        first = next(iterator, None)
        if first is None:
            raise ProductImportTerminalError(
                "The file is empty."
            )

        layout = _header_layout(first)
        headers = [
            header
            for _, header in layout
        ]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            raw = _row_from_layout(
                values,
                layout,
            )
            if not any(
                value is not None
                for value in raw.values()
            ):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise ProductImportTerminalError(
                    f"The import exceeds the {MAX_IMPORT_ROWS:,}-row safety limit."
                )
            rows.append(raw)

        if not rows:
            raise ProductImportTerminalError(
                "The file contains no product rows."
            )

        return headers, rows
    finally:
        workbook.close()


def parse_source(
    file_name: str,
    payload: bytes,
) -> tuple[list[str], list[dict[str, Any]]]:
    suffix = (
        file_name.lower().rsplit(".", 1)[-1]
        if "." in file_name
        else ""
    )
    if suffix == "csv":
        return _parse_csv(payload)
    if suffix == "xlsx":
        return _parse_xlsx(payload)
    raise ProductImportTerminalError(
        "Only CSV and XLSX files are supported."
    )


