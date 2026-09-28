from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from domains.simple_products.imports.domain import (
    AR_IMPORT_LOCALE,
    CANONICAL_IMPORT_FIELDS,
    EN_IMPORT_LOCALE,
    IMPORT_TRACKING_DEFAULT_SENTINEL,
    ImportLocalePack,
    WANASAH_TEMPLATE_MARKER,
    WANASAH_TEMPLATE_META_SHEET,
)


MAX_TEMPLATE_ROWS = 50_000
_TRACKING_FIELDS = (
    "lot_control_mode",
    "expiry_control_mode",
)
_BARCODE_FIELDS = (
    "unit_barcode",
    "package_barcode",
)
_PACKAGE_CODES = (
    "CARTON",
    "CASE",
    "PACK",
    "BAG",
    "SACK",
    "TRAY",
    "CRATE",
    "BUNDLE",
    "PALLET",
)


def _locale_pack(locale: str) -> ImportLocalePack:
    return (
        EN_IMPORT_LOCALE
        if str(locale).lower().startswith("en")
        else AR_IMPORT_LOCALE
    )


def _preferred_tracking_label(
    pack: ImportLocalePack,
    canonical: str,
) -> str:
    for label, code in (
        pack.tracking_value_aliases.items()
    ):
        if code == canonical:
            return label
    raise ValueError(
        f"Missing tracking label for {canonical!r} "
        f"in locale {pack.locale!r}."
    )


def _preferred_package_label(
    pack: ImportLocalePack,
    canonical: str,
) -> str:
    for label, code in (
        pack.package_value_aliases.items()
    ):
        if code == canonical:
            return label
    raise ValueError(
        f"Missing package label for {canonical!r} "
        f"in locale {pack.locale!r}."
    )


def _template_copy(
    locale: str,
) -> dict[str, str]:
    if locale == "en":
        return {
            "sheet": "Products",
            "default_note": (
                "Leave blank, or choose Use default, to use the "
                "tracking default shown in Wanasah when you upload "
                "this file. Choose another value only for exceptions."
            ),
            "validation_error": (
                "Choose a value from the list, or leave the cell blank."
            ),
            "validation_title": "Product tracking",
            "barcode_validation_title": "Barcode text",
            "barcode_validation_error": (
                "Enter the complete barcode as text. "
                "Do not use numbers, formulas, or scientific notation."
            ),
            "package_note": (
                "Choose the outer package type. Leave blank if the product "
                "is sold only as individual units."
            ),
            "units_note": (
                "Required only when an outer package is selected. "
                "Enter a whole number from 2 to 1,000,000."
            ),
            "package_price_note": (
                "Optional. Use only when an outer package is selected. "
                "Enter a positive number."
            ),
            "unit_price_note": (
                "Enter a positive number when provided. At least one of "
                "Package Price or Unit Price must be supplied per product."
            ),
            "unit_barcode_note": (
                "Optional. Keep the complete unit barcode as text, "
                "especially when it starts with zero."
            ),
            "package_barcode_note": (
                "Optional and only valid when an outer package is selected. "
                "Keep the complete barcode as text."
            ),
            "package_validation_title": "Package type",
            "units_validation_title": "Units per package",
            "price_validation_title": "Price",
        }
    return {
        "sheet": "المنتجات",
        "default_note": (
            "اترك الخانة فارغة، أو اختر «استخدام الافتراضي»، "
            "ليستخدم المنتج إعداد التتبع الافتراضي الظاهر في وناسة "
            "وقت رفع الملف. اختر قيمة أخرى فقط للاستثناءات."
        ),
        "validation_error": (
            "اختر قيمة من القائمة، أو اترك الخانة فارغة."
        ),
        "validation_title": "تتبع المنتج",
        "barcode_validation_title": "الباركود كنص",
        "barcode_validation_error": (
            "أدخل الباركود كاملاً كنص. لا تستخدم رقماً أو معادلة أو صيغة علمية."
        ),
        "package_note": (
            "اختر نوع العبوة الخارجية. اترك الخانة فارغة إذا كان المنتج "
            "يباع كوحدات مفردة فقط."
        ),
        "units_note": (
            "مطلوب فقط عند اختيار عبوة خارجية. أدخل عدداً صحيحاً "
            "من 2 إلى 1,000,000."
        ),
        "package_price_note": (
            "اختياري، ويستخدم فقط عند وجود عبوة خارجية. أدخل رقماً موجباً."
        ),
        "unit_price_note": (
            "أدخل رقماً موجباً عند تعبئته. يجب توفير سعر العبوة أو سعر الوحدة "
            "على الأقل لكل منتج."
        ),
        "unit_barcode_note": (
            "اختياري. احتفظ بالباركود كاملاً كنص، خصوصاً إذا بدأ بصفر."
        ),
        "package_barcode_note": (
            "اختياري ويقبل فقط عند وجود عبوة خارجية. احتفظ بالباركود كاملاً كنص."
        ),
        "package_validation_title": "نوع العبوة",
        "units_validation_title": "عدد الوحدات",
        "price_validation_title": "السعر",
    }


def build_product_import_template(
    *,
    locale: str,
) -> bytes:
    pack = _locale_pack(locale)
    copy = _template_copy(pack.locale)

    headers = [
        pack.header_aliases[field][0]
        for field in CANONICAL_IMPORT_FIELDS
    ]
    tracking_values = [
        _preferred_tracking_label(
            pack,
            IMPORT_TRACKING_DEFAULT_SENTINEL,
        ),
        _preferred_tracking_label(
            pack,
            "NONE",
        ),
        _preferred_tracking_label(
            pack,
            "OPTIONAL",
        ),
        _preferred_tracking_label(
            pack,
            "REQUIRED",
        ),
    ]
    package_values = [
        _preferred_package_label(
            pack,
            code,
        )
        for code in _PACKAGE_CODES
    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = copy["sheet"]
    sheet.freeze_panes = "A2"
    sheet.sheet_view.rightToLeft = (
        pack.locale == "ar"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="475569",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11,
    )
    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column_index,
            value=header,
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    sheet.auto_filter.ref = (
        f"A1:J{MAX_TEMPLATE_ROWS + 1}"
    )
    sheet.row_dimensions[1].height = 30

    widths = (
        30,
        20,
        16,
        18,
        14,
        14,
        20,
        20,
        16,
        16,
    )
    for index, width in enumerate(
        widths,
        start=1,
    ):
        sheet.column_dimensions[
            chr(64 + index)
        ].width = width

    package_column = "C"
    units_column = "D"
    package_price_column = "E"
    unit_price_column = "F"
    package_barcode_column = "H"

    package_validation = DataValidation(
        type="list",
        formula1="='_wanasah_lists'!$B$1:$B$9",
        allow_blank=True,
    )
    package_validation.errorTitle = copy[
        "package_validation_title"
    ]
    package_validation.error = copy[
        "validation_error"
    ]
    package_validation.promptTitle = copy[
        "package_validation_title"
    ]
    package_validation.prompt = copy[
        "package_note"
    ]
    package_validation.showErrorMessage = True
    package_validation.showInputMessage = True
    sheet.add_data_validation(
        package_validation
    )
    package_validation.add(
        f"{package_column}2:"
        f"{package_column}{MAX_TEMPLATE_ROWS + 1}"
    )

    units_validation = DataValidation(
        type="custom",
        formula1=(
            '=IF(C2="",D2="",'
            'AND(ISNUMBER(D2),D2=INT(D2),'
            'D2>=2,D2<=1000000))'
        ),
        allow_blank=False,
    )
    units_validation.errorTitle = copy[
        "units_validation_title"
    ]
    units_validation.error = copy[
        "units_note"
    ]
    units_validation.promptTitle = copy[
        "units_validation_title"
    ]
    units_validation.prompt = copy[
        "units_note"
    ]
    units_validation.showErrorMessage = True
    units_validation.showInputMessage = True
    sheet.add_data_validation(
        units_validation
    )
    units_validation.add(
        f"{units_column}2:"
        f"{units_column}{MAX_TEMPLATE_ROWS + 1}"
    )

    package_price_validation = DataValidation(
        type="custom",
        formula1=(
            '=OR(E2="",AND(C2<>"",'
            'ISNUMBER(E2),E2>0))'
        ),
        allow_blank=True,
    )
    package_price_validation.errorTitle = copy[
        "price_validation_title"
    ]
    package_price_validation.error = copy[
        "package_price_note"
    ]
    package_price_validation.promptTitle = copy[
        "price_validation_title"
    ]
    package_price_validation.prompt = copy[
        "package_price_note"
    ]
    package_price_validation.showErrorMessage = True
    package_price_validation.showInputMessage = True
    sheet.add_data_validation(
        package_price_validation
    )
    package_price_validation.add(
        f"{package_price_column}2:"
        f"{package_price_column}{MAX_TEMPLATE_ROWS + 1}"
    )

    unit_price_validation = DataValidation(
        type="custom",
        formula1=(
            '=OR(F2="",AND(ISNUMBER(F2),F2>0))'
        ),
        allow_blank=True,
    )
    unit_price_validation.errorTitle = copy[
        "price_validation_title"
    ]
    unit_price_validation.error = copy[
        "unit_price_note"
    ]
    unit_price_validation.promptTitle = copy[
        "price_validation_title"
    ]
    unit_price_validation.prompt = copy[
        "unit_price_note"
    ]
    unit_price_validation.showErrorMessage = True
    unit_price_validation.showInputMessage = True
    sheet.add_data_validation(
        unit_price_validation
    )
    unit_price_validation.add(
        f"{unit_price_column}2:"
        f"{unit_price_column}{MAX_TEMPLATE_ROWS + 1}"
    )

    sheet.column_dimensions[
        units_column
    ].number_format = "0"
    sheet.column_dimensions[
        package_price_column
    ].number_format = "0.000"
    sheet.column_dimensions[
        unit_price_column
    ].number_format = "0.000"

    for field in _BARCODE_FIELDS:
        column_index = (
            CANONICAL_IMPORT_FIELDS.index(
                field
            )
            + 1
        )
        column_letter = chr(
            64 + column_index
        )
        sheet.column_dimensions[
            column_letter
        ].number_format = "@"

        barcode_formula = (
            (
                f'=OR({column_letter}2="",'
                f'AND(C2<>"",ISTEXT({column_letter}2)))'
            )
            if column_letter
            == package_barcode_column
            else (
                f'=OR({column_letter}2="",'
                f'ISTEXT({column_letter}2))'
            )
        )
        barcode_validation = DataValidation(
            type="custom",
            formula1=barcode_formula,
            allow_blank=True,
        )
        barcode_validation.errorTitle = (
            copy[
                "barcode_validation_title"
            ]
        )
        barcode_validation.error = copy[
            "barcode_validation_error"
        ]
        barcode_validation.showErrorMessage = True
        sheet.add_data_validation(
            barcode_validation
        )
        barcode_validation.add(
            f"{column_letter}2:"
            f"{column_letter}{MAX_TEMPLATE_ROWS + 1}"
        )

    meta_sheet = workbook.create_sheet(
        WANASAH_TEMPLATE_META_SHEET
    )
    meta_sheet["A1"] = (
        WANASAH_TEMPLATE_MARKER
    )
    meta_sheet["A2"] = sheet.title
    meta_sheet.sheet_state = (
        "veryHidden"
    )

    list_sheet = workbook.create_sheet(
        "_wanasah_lists"
    )
    for row_index, value in enumerate(
        tracking_values,
        start=1,
    ):
        list_sheet.cell(
            row=row_index,
            column=1,
            value=value,
        )
    for row_index, value in enumerate(
        package_values,
        start=1,
    ):
        list_sheet.cell(
            row=row_index,
            column=2,
            value=value,
        )
    list_sheet.sheet_state = "hidden"

    formula = (
        "='_wanasah_lists'!$A$1:$A$4"
    )
    for field in _TRACKING_FIELDS:
        column_index = (
            CANONICAL_IMPORT_FIELDS.index(
                field
            )
            + 1
        )
        column_letter = chr(
            64 + column_index
        )
        validation = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
        )
        validation.errorTitle = (
            copy[
                "validation_title"
            ]
        )
        validation.error = copy[
            "validation_error"
        ]
        validation.promptTitle = (
            copy[
                "validation_title"
            ]
        )
        validation.prompt = copy[
            "default_note"
        ]
        validation.showErrorMessage = True
        validation.showInputMessage = True

        sheet.add_data_validation(
            validation
        )
        validation.add(
            f"{column_letter}2:"
            f"{column_letter}{MAX_TEMPLATE_ROWS + 1}"
        )

        sheet.cell(
            row=1,
            column=column_index,
        ).comment = Comment(
            copy["default_note"],
            "Wanasah",
        )

    header_notes = {
        "name": (
            "Required."
            if pack.locale
            == "en"
            else "مطلوب."
        ),
        "family": (
            "Optional product family."
            if pack.locale
            == "en"
            else "عائلة المنتج اختيارية."
        ),
        "package_uom":
            copy["package_note"],
        "units_per_package":
            copy["units_note"],
        "package_price":
            copy["package_price_note"],
        "unit_price":
            copy["unit_price_note"],
        "unit_barcode":
            copy["unit_barcode_note"],
        "package_barcode":
            copy["package_barcode_note"],
        "lot_control_mode":
            copy["default_note"],
        "expiry_control_mode":
            copy["default_note"],
    }
    for field, note in (
        header_notes.items()
    ):
        column_index = (
            CANONICAL_IMPORT_FIELDS.index(
                field
            )
            + 1
        )
        sheet.cell(
            row=1,
            column=column_index,
        ).comment = Comment(
            note,
            "Wanasah",
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
