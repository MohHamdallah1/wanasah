"""One-time opening inventory valuation for stock that predates costing."""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
import hashlib
import json
from typing import Any
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from models import (
    Company,
    InventoryBalance,
    InventoryCostEvent,
    InventoryCostLayer,
    InventoryCostPolicy,
    InventoryCostState,
    InventoryOpeningValuation,
    InventoryOpeningValuationLine,
    ProductVariant,
    SystemAuditLog,
    UOM,
    utc_now,
)
from domains.inventory_costing.service import (
    COST_METHODS,
    CostingError,
    _policy_guard,
    acquire_inventory_cost_bootstrap_guard,
    money6,
)


@dataclass(frozen=True)
class OpeningBatchSnapshot:
    product_variant_id: int
    batch_id: int
    quantity: Decimal


@dataclass(frozen=True)
class OpeningVariantSnapshot:
    product_variant_id: int
    sku: str
    name: str
    base_uom_id: int
    base_uom_code: str
    base_uom_name: str
    quantity: Decimal
    batches: tuple[OpeningBatchSnapshot, ...]


@dataclass(frozen=True)
class OpeningSnapshot:
    hash: str
    variants: tuple[OpeningVariantSnapshot, ...]
    total_quantity: Decimal


def _snapshot_hash(batch_rows: list[tuple[int, int, Decimal]]) -> str:
    payload = [
        [int(variant_id), int(batch_id), format(Decimal(quantity), "f")]
        for variant_id, batch_id, quantity in batch_rows
    ]
    raw = json.dumps(payload, separators=(",", ":"), ensure_ascii=True)
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


async def load_opening_snapshot(
    db: AsyncSession,
    *,
    company_id: int,
) -> OpeningSnapshot:
    """Read the exact company-wide positive physical quantity snapshot."""
    rows = (
        await db.execute(
            select(
                InventoryBalance.product_variant_id,
                InventoryBalance.batch_id,
                func.sum(InventoryBalance.on_hand_quantity).label("quantity"),
            )
            .where(
                InventoryBalance.company_id == int(company_id),
                InventoryBalance.on_hand_quantity > 0,
            )
            .group_by(
                InventoryBalance.product_variant_id,
                InventoryBalance.batch_id,
            )
            .having(func.sum(InventoryBalance.on_hand_quantity) > 0)
            .order_by(
                InventoryBalance.product_variant_id.asc(),
                InventoryBalance.batch_id.asc(),
            )
        )
    ).all()
    batch_rows = [
        (int(row.product_variant_id), int(row.batch_id), Decimal(row.quantity))
        for row in rows
    ]
    if not batch_rows:
        return OpeningSnapshot(
            hash=_snapshot_hash([]),
            variants=(),
            total_quantity=Decimal("0"),
        )

    variant_ids = sorted({row[0] for row in batch_rows})
    metadata_rows = (
        await db.execute(
            select(
                ProductVariant.id,
                ProductVariant.sku,
                ProductVariant.name,
                ProductVariant.base_uom_id,
                UOM.code,
                UOM.name.label("uom_name"),
            )
            .join(UOM, UOM.id == ProductVariant.base_uom_id)
            .where(
                ProductVariant.company_id == int(company_id),
                ProductVariant.id.in_(variant_ids),
            )
            .order_by(ProductVariant.id.asc())
        )
    ).all()
    metadata = {int(row.id): row for row in metadata_rows}
    if set(metadata) != set(variant_ids):
        raise CostingError(
            "OPENING_VALUATION_VARIANT_MISSING",
            "One or more stocked product variants could not be resolved.",
        )

    by_variant: dict[int, list[OpeningBatchSnapshot]] = {}
    for variant_id, batch_id, quantity in batch_rows:
        by_variant.setdefault(variant_id, []).append(
            OpeningBatchSnapshot(
                product_variant_id=variant_id,
                batch_id=batch_id,
                quantity=quantity,
            )
        )

    variants: list[OpeningVariantSnapshot] = []
    total_quantity = Decimal("0")
    for variant_id in variant_ids:
        row = metadata[variant_id]
        batches = tuple(by_variant[variant_id])
        quantity = sum((batch.quantity for batch in batches), Decimal("0"))
        total_quantity += quantity
        variants.append(
            OpeningVariantSnapshot(
                product_variant_id=variant_id,
                sku=str(row.sku),
                name=str(row.name),
                base_uom_id=int(row.base_uom_id),
                base_uom_code=str(row.code),
                base_uom_name=str(row.uom_name),
                quantity=quantity,
                batches=batches,
            )
        )
    return OpeningSnapshot(
        hash=_snapshot_hash(batch_rows),
        variants=tuple(variants),
        total_quantity=total_quantity,
    )


def _cost_map(
    raw_lines: list[dict[str, Any]],
) -> dict[str, Decimal]:
    result: dict[str, Decimal] = {}
    for index, raw in enumerate(raw_lines):
        sku = str(raw.get("sku") or "").strip()
        if not sku or len(sku) > 100:
            raise CostingError(
                "OPENING_VALUATION_SKU_INVALID",
                "Opening valuation SKU is invalid.",
                status_code=422,
                context={"row": index + 1},
            )
        if sku in result:
            raise CostingError(
                "OPENING_VALUATION_DUPLICATE_SKU",
                "Opening valuation contains a duplicate SKU.",
                status_code=422,
                context={"sku": sku},
            )
        result[sku] = money6(raw.get("unit_cost"), "unit_cost")
    return result


def opening_preview_payload(
    *,
    snapshot: OpeningSnapshot,
    raw_lines: list[dict[str, Any]],
    currency_code: str,
) -> dict[str, Any]:
    costs = _cost_map(raw_lines)
    expected = {variant.sku for variant in snapshot.variants}
    supplied = set(costs)
    missing = sorted(expected - supplied)
    extra = sorted(supplied - expected)
    if missing or extra:
        raise CostingError(
            "OPENING_VALUATION_COVERAGE_MISMATCH",
            "Opening valuation must cover every currently stocked SKU exactly once.",
            status_code=409,
            context={
                "missing_count": len(missing),
                "extra_count": len(extra),
                "missing_skus": missing[:20],
                "extra_skus": extra[:20],
            },
        )
    total_value = Decimal("0")
    zero_count = 0
    items: list[dict[str, Any]] = []
    for variant in snapshot.variants:
        cost = costs[variant.sku]
        if cost == 0:
            zero_count += 1
        value = sum(
            (
                money6(batch.quantity * cost, "batch_total_value")
                for batch in variant.batches
            ),
            Decimal("0"),
        )
        value = money6(value, "total_value")
        total_value += value
        items.append(
            {
                "sku": variant.sku,
                "name": variant.name,
                "base_uom_code": variant.base_uom_code,
                "base_uom_name": variant.base_uom_name,
                "quantity": format(variant.quantity, "f"),
                "unit_cost": format(cost, "f"),
                "total_value": format(value, "f"),
            }
        )
    return {
        "snapshot_hash": snapshot.hash,
        "currency_code": str(currency_code).upper(),
        "line_count": len(snapshot.variants),
        "total_quantity": format(snapshot.total_quantity, "f"),
        "total_value": format(money6(total_value, "total_value"), "f"),
        "zero_cost_count": zero_count,
        "items": items,
    }


async def _ensure_opening_history_is_empty(
    db: AsyncSession,
    *,
    company_id: int,
) -> None:
    if await db.scalar(
        select(func.count(InventoryOpeningValuation.id)).where(
            InventoryOpeningValuation.company_id == int(company_id)
        )
    ):
        raise CostingError(
            "OPENING_VALUATION_ALREADY_POSTED",
            "Opening inventory valuation was already posted.",
        )
    for model in (InventoryCostState, InventoryCostEvent, InventoryCostLayer):
        if await db.scalar(
            select(func.count(model.id)).where(
                model.company_id == int(company_id)
            )
        ):
            raise CostingError(
                "OPENING_VALUATION_COST_HISTORY_EXISTS",
                "Cost history already exists; opening valuation cannot be posted.",
            )


async def load_opening_valuation_context(
    db: AsyncSession,
    *,
    company_id: int,
) -> tuple[Company, InventoryCostPolicy, OpeningSnapshot]:
    company = await db.scalar(
        select(Company).where(Company.id == int(company_id))
    )
    if company is None:
        raise CostingError("COMPANY_NOT_FOUND", "Company was not found.", status_code=404)
    policy = await db.scalar(
        select(InventoryCostPolicy).where(
            InventoryCostPolicy.company_id == int(company_id)
        )
    )
    if policy is None or policy.selected_at is None or policy.selected_by is None:
        raise CostingError(
            "INVENTORY_COST_POLICY_SELECTION_REQUIRED",
            "Choose and save a costing method before opening valuation.",
        )
    if policy.is_active or policy.locked_at is not None:
        raise CostingError(
            "INVENTORY_COST_POLICY_LOCKED",
            "Inventory costing is already active.",
        )
    await _ensure_opening_history_is_empty(
        db,
        company_id=int(company_id),
    )
    snapshot = await load_opening_snapshot(db, company_id=int(company_id))
    if not snapshot.variants:
        raise CostingError(
            "OPENING_VALUATION_NOT_REQUIRED",
            "There is no positive opening inventory to value.",
        )
    return company, policy, snapshot


async def preview_opening_valuation(
    db: AsyncSession,
    *,
    company_id: int,
    raw_lines: list[dict[str, Any]],
    expected_snapshot_hash: str | None = None,
) -> dict[str, Any]:
    company, _policy, snapshot = await load_opening_valuation_context(
        db,
        company_id=int(company_id),
    )
    if (
        expected_snapshot_hash is not None
        and snapshot.hash != str(expected_snapshot_hash).lower()
    ):
        raise CostingError(
            "OPENING_VALUATION_TEMPLATE_STALE",
            "Inventory changed after this template was downloaded.",
            context={
                "template_snapshot_hash": str(expected_snapshot_hash).lower(),
                "current_snapshot_hash": snapshot.hash,
            },
        )
    return opening_preview_payload(
        snapshot=snapshot,
        raw_lines=raw_lines,
        currency_code=str(company.currency_code),
    )



OPENING_TEMPLATE_MAX_BYTES = 5 * 1024 * 1024
OPENING_TEMPLATE_MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
OPENING_TEMPLATE_MAX_ARCHIVE_ENTRIES = 1000
OPENING_TEMPLATE_SHEET = "Opening Valuation"
OPENING_TEMPLATE_META_SHEET = "_wanasah_meta"


def build_opening_valuation_template(
    *,
    snapshot: OpeningSnapshot,
    currency_code: str,
    locale: str = "ar",
) -> bytes:
    arabic = str(locale).lower().startswith("ar")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = OPENING_TEMPLATE_SHEET
    headers = (
        [
            "SKU",
            "المنتج",
            "وحدة الأساس",
            "الكمية الحالية",
            f"تكلفة الوحدة الافتتاحية ({str(currency_code).upper()})",
        ]
        if arabic
        else [
            "SKU",
            "Product",
            "Base UOM",
            "Current quantity",
            f"Opening unit cost ({str(currency_code).upper()})",
        ]
    )
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")
    for variant in snapshot.variants:
        sheet.append(
            [
                variant.sku,
                variant.name,
                f"{variant.base_uom_code} — {variant.base_uom_name}",
                float(variant.quantity),
                None,
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
    widths = {"A": 22, "B": 34, "C": 24, "D": 18, "E": 28}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, max_col=5):
        for cell in row[:4]:
            cell.protection = Protection(locked=True)
        row[4].protection = Protection(locked=False)
        row[3].number_format = "0.######"
        row[4].number_format = "0.000000"
    sheet.protection.sheet = True
    sheet.protection.enable()
    sheet.protection.password = "wanasah"

    meta = workbook.create_sheet(OPENING_TEMPLATE_META_SHEET)
    meta.sheet_state = "veryHidden"
    meta.append(["contract", "opening-valuation-v1"])
    meta.append(["snapshot_hash", snapshot.hash])
    meta.append(["currency_code", str(currency_code).upper()])
    meta.append(["line_count", len(snapshot.variants)])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_opening_valuation_template(
    payload: bytes,
) -> tuple[str, list[dict[str, Any]]]:
    if not payload or len(payload) > OPENING_TEMPLATE_MAX_BYTES:
        raise CostingError(
            "OPENING_VALUATION_FILE_INVALID",
            "Opening valuation file is empty or too large.",
            status_code=422,
        )
    try:
        with ZipFile(BytesIO(payload)) as archive:
            infos = archive.infolist()
            if (
                len(infos) > OPENING_TEMPLATE_MAX_ARCHIVE_ENTRIES
                or sum(info.file_size for info in infos)
                > OPENING_TEMPLATE_MAX_UNCOMPRESSED_BYTES
            ):
                raise CostingError(
                    "OPENING_VALUATION_FILE_INVALID",
                    "Opening valuation workbook is too large after decompression.",
                    status_code=422,
                )
            for info in infos:
                normalized = info.filename.replace("\\", "/")
                if normalized.startswith("/") or ".." in normalized.split("/"):
                    raise CostingError(
                        "OPENING_VALUATION_FILE_INVALID",
                        "Opening valuation workbook contains an unsafe archive path.",
                        status_code=422,
                    )
    except BadZipFile as exc:
        raise CostingError(
            "OPENING_VALUATION_FILE_INVALID",
            "Opening valuation file is not a valid XLSX workbook.",
            status_code=422,
        ) from exc

    try:
        workbook = load_workbook(
            BytesIO(payload),
            read_only=True,
            data_only=False,
        )
    except Exception as exc:
        raise CostingError(
            "OPENING_VALUATION_FILE_INVALID",
            "Opening valuation file could not be read.",
            status_code=422,
        ) from exc
    try:
        if (
            OPENING_TEMPLATE_SHEET not in workbook.sheetnames
            or OPENING_TEMPLATE_META_SHEET not in workbook.sheetnames
        ):
            raise CostingError(
                "OPENING_VALUATION_FILE_CONTRACT_INVALID",
                "Use the opening valuation template downloaded from Wanasah.",
                status_code=422,
            )
        meta = workbook[OPENING_TEMPLATE_META_SHEET]
        metadata = {
            str(row[0].value): row[1].value
            for row in meta.iter_rows(min_row=1, max_col=2)
            if row[0].value is not None
        }
        if metadata.get("contract") != "opening-valuation-v1":
            raise CostingError(
                "OPENING_VALUATION_FILE_CONTRACT_INVALID",
                "Opening valuation workbook contract is invalid.",
                status_code=422,
            )
        snapshot_hash = str(metadata.get("snapshot_hash") or "").lower()
        if (
            len(snapshot_hash) != 64
            or any(ch not in "0123456789abcdef" for ch in snapshot_hash)
        ):
            raise CostingError(
                "OPENING_VALUATION_FILE_CONTRACT_INVALID",
                "Opening valuation workbook snapshot is invalid.",
                status_code=422,
            )

        sheet = workbook[OPENING_TEMPLATE_SHEET]
        if sheet.max_row > 10001:
            raise CostingError(
                "OPENING_VALUATION_FILE_TOO_MANY_ROWS",
                "Opening valuation workbook has too many rows.",
                status_code=422,
            )
        lines: list[dict[str, Any]] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, max_col=5, values_only=True),
            start=2,
        ):
            sku = row[0]
            unit_cost = row[4]
            if sku is None and unit_cost is None:
                continue
            sku_text = str(sku or "").strip()
            if not sku_text:
                raise CostingError(
                    "OPENING_VALUATION_SKU_INVALID",
                    "Opening valuation row is missing SKU.",
                    status_code=422,
                    context={"row": row_number},
                )
            if unit_cost is None or (
                isinstance(unit_cost, str) and not unit_cost.strip()
            ):
                raise CostingError(
                    "OPENING_VALUATION_COST_REQUIRED",
                    "Opening unit cost is required for every stocked SKU.",
                    status_code=422,
                    context={"row": row_number, "sku": sku_text},
                )
            if isinstance(unit_cost, str) and unit_cost.lstrip().startswith("="):
                raise CostingError(
                    "OPENING_VALUATION_FORMULA_NOT_ALLOWED",
                    "Opening unit cost must be a value, not a formula.",
                    status_code=422,
                    context={"row": row_number, "sku": sku_text},
                )
            lines.append({"sku": sku_text, "unit_cost": unit_cost})
        if not lines:
            raise CostingError(
                "OPENING_VALUATION_FILE_EMPTY",
                "Opening valuation workbook has no cost rows.",
                status_code=422,
            )
        return snapshot_hash, lines
    finally:
        workbook.close()


async def finalize_opening_valuation(
    db: AsyncSession,
    *,
    company_id: int,
    actor_id: int,
    request_id: UUID,
    expected_snapshot_hash: str,
    raw_lines: list[dict[str, Any]],
    confirm_zero_costs: bool,
) -> dict[str, Any]:
    # Lock order is policy -> company bootstrap. Supplier first activation follows
    # the same policy-before-bootstrap order through activate + movement posting.
    await _policy_guard(db, int(company_id))
    policy = await db.scalar(
        select(InventoryCostPolicy)
        .where(InventoryCostPolicy.company_id == int(company_id))
        .with_for_update()
    )
    if policy is None or policy.selected_at is None or policy.selected_by is None:
        raise CostingError(
            "INVENTORY_COST_POLICY_SELECTION_REQUIRED",
            "Choose and save a costing method before opening valuation.",
        )
    if str(policy.method).upper() not in COST_METHODS:
        raise CostingError(
            "INVENTORY_COST_METHOD_INVALID",
            "Inventory costing policy is invalid.",
        )
    if policy.is_active or policy.locked_at is not None:
        raise CostingError(
            "INVENTORY_COST_POLICY_LOCKED",
            "Inventory costing is already active.",
        )

    await acquire_inventory_cost_bootstrap_guard(
        db,
        company_id=int(company_id),
        exclusive=True,
    )

    await _ensure_opening_history_is_empty(
        db,
        company_id=int(company_id),
    )

    company = await db.scalar(
        select(Company).where(Company.id == int(company_id))
    )
    if company is None:
        raise CostingError("COMPANY_NOT_FOUND", "Company was not found.", status_code=404)

    snapshot = await load_opening_snapshot(db, company_id=int(company_id))
    if not snapshot.variants:
        raise CostingError(
            "OPENING_VALUATION_NOT_REQUIRED",
            "There is no positive opening inventory to value.",
        )
    if snapshot.hash != str(expected_snapshot_hash).lower():
        raise CostingError(
            "OPENING_VALUATION_SNAPSHOT_STALE",
            "Inventory changed after preview; refresh opening valuation and try again.",
            context={
                "expected_snapshot_hash": str(expected_snapshot_hash).lower(),
                "current_snapshot_hash": snapshot.hash,
            },
        )
    preview = opening_preview_payload(
        snapshot=snapshot,
        raw_lines=raw_lines,
        currency_code=str(company.currency_code),
    )
    if preview["zero_cost_count"] and not confirm_zero_costs:
        raise CostingError(
            "OPENING_VALUATION_ZERO_COST_CONFIRMATION_REQUIRED",
            "One or more opening costs are zero and require explicit confirmation.",
            status_code=422,
            context={"zero_cost_count": preview["zero_cost_count"]},
        )

    costs = _cost_map(raw_lines)
    method = str(policy.method).upper()
    now = utc_now()
    header = InventoryOpeningValuation(
        company_id=int(company_id),
        request_id=request_id,
        method=method,
        currency_code=str(company.currency_code).upper(),
        inventory_snapshot_hash=snapshot.hash,
        line_count=len(snapshot.variants),
        total_quantity=Decimal(preview["total_quantity"]),
        total_value=Decimal(preview["total_value"]),
        posted_by=int(actor_id),
        posted_at=now,
    )
    db.add(header)
    await db.flush()

    for variant in snapshot.variants:
        unit_cost = costs[variant.sku]
        variant_value = money6(
            sum(
                (
                    money6(batch.quantity * unit_cost, "batch_total_value")
                    for batch in variant.batches
                ),
                Decimal("0"),
            ),
            "total_value",
        )
        line = InventoryOpeningValuationLine(
            company_id=int(company_id),
            opening_valuation_id=int(header.id),
            product_variant_id=variant.product_variant_id,
            quantity=variant.quantity,
            unit_cost=unit_cost,
            total_value=variant_value,
        )
        db.add(line)
        await db.flush()

        db.add(
            InventoryCostState(
                company_id=int(company_id),
                product_variant_id=variant.product_variant_id,
                quantity=variant.quantity,
                inventory_value=variant_value,
                average_unit_cost=unit_cost,
                version=1,
            )
        )

        quantity_before = Decimal("0")
        value_before = Decimal("0")
        for batch in variant.batches:
            event_value = money6(batch.quantity * unit_cost, "total_cost")
            event = InventoryCostEvent(
                company_id=int(company_id),
                inventory_movement_id=None,
                opening_valuation_line_id=int(line.id),
                product_variant_id=variant.product_variant_id,
                batch_id=batch.batch_id,
                method=method,
                event_type="OPENING_BALANCE",
                cost_basis="OPENING_VALUATION",
                quantity=batch.quantity,
                unit_cost=unit_cost,
                total_cost=event_value,
                quantity_before=quantity_before,
                quantity_after=quantity_before + batch.quantity,
                value_before=value_before,
                value_after=money6(value_before + event_value, "value_after"),
                input_uom_id=None,
                input_quantity=None,
                input_unit_cost=None,
                reversal_of_cost_event_id=None,
                created_at=now,
            )
            db.add(event)
            await db.flush()
            if method == "FIFO":
                db.add(
                    InventoryCostLayer(
                        company_id=int(company_id),
                        product_variant_id=variant.product_variant_id,
                        batch_id=batch.batch_id,
                        source_cost_event_id=int(event.id),
                        original_quantity=batch.quantity,
                        remaining_quantity=batch.quantity,
                        unit_cost=unit_cost,
                        original_value=event_value,
                        remaining_value=event_value,
                        version=1,
                        created_at=now,
                        updated_at=now,
                    )
                )
            quantity_before += batch.quantity
            value_before = money6(value_before + event_value, "value_before")

    policy.is_active = True
    policy.locked_at = now
    policy.version += 1
    policy.updated_by = int(actor_id)
    policy.updated_at = now

    db.add(
        SystemAuditLog(
            company_id=int(company_id),
            admin_id=int(actor_id),
            target_id=f"InventoryOpeningValuation_{header.id}",
            action_type="INVENTORY_OPENING_VALUATION_POSTED",
            old_value=None,
            new_value=json.dumps(
                {
                    "request_id": str(request_id),
                    "method": method,
                    "currency_code": str(company.currency_code).upper(),
                    "snapshot_hash": snapshot.hash,
                    "line_count": len(snapshot.variants),
                    "total_quantity": preview["total_quantity"],
                    "total_value": preview["total_value"],
                    "zero_cost_count": preview["zero_cost_count"],
                },
                separators=(",", ":"),
                sort_keys=True,
            ),
        )
    )
    await db.flush()

    return {
        "message": "OPENING_VALUATION_POSTED",
        "opening_valuation_id": int(header.id),
        "method": method,
        "currency_code": str(company.currency_code).upper(),
        "snapshot_hash": snapshot.hash,
        "line_count": len(snapshot.variants),
        "total_quantity": preview["total_quantity"],
        "total_value": preview["total_value"],
        "zero_cost_count": preview["zero_cost_count"],
        "policy_version": int(policy.version),
    }
